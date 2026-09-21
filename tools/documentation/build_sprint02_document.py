from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "sprint-02" / "assets"
OUT = ROOT / "output" / "docx"
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "Grupo_04_RotaCerta_Sprints_01_02.docx"

NAVY = "173B67"
PALE_BLUE = "EAF2FF"
PALE_GRAY = "F5F7FA"
LIGHT_BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(88, 96, 110)


def set_font(run, name="Arial", size=11, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), LIGHT_BORDER)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        set_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for r_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if r_idx % 2:
                set_cell_shading(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx == 0 and len(headers) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_font(run, size=font_size)
    if widths:
        set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc, items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{idx}. {item}")
        set_font(r)


def add_figure(doc, filename, caption, width=6.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_font(r, size=9, italic=True, color=GRAY)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_font(run, size=9, color=GRAY)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(23)
    title_style.font.bold = True
    title_style.font.color.rgb = BLACK
    title_ppr = title_style.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("RotaCerta  |  Sprints 01 e 02  |  Página ")
    set_font(r, size=9, color=GRAY)
    add_page_number(footer)

    doc.core_properties.title = "RotaCerta Documento Técnico Cumulativo Sprints 01 e 02"
    doc.core_properties.subject = "Planejamento, arquitetura e modelagem do sistema"
    doc.core_properties.author = "Equipe RotaCerta"
    doc.core_properties.keywords = "RotaCerta, Sprint 01, Sprint 02, arquitetura, modelagem"
    return doc


def build():
    doc = setup_document()

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("UNINASSAU")
    set_font(r, size=14, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(68)
    r = p.add_run("Fábrica de Software e Tópicos Avançados em Ciência da Computação\nTurma 8NA  |  Semestre 2026.2")
    set_font(r, size=11)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.add_run("RotaCerta")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Documento Técnico Cumulativo das Sprints 01 e 02")
    set_font(r, size=17, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(64)
    r = p.add_run("Grupo 04")
    set_font(r, size=14, bold=True)

    for text in ("Álvaro Jordão  |  01748200", "Arthur Sales  |  01593811", "Vinícius Trigueiro  |  01794959"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text)
        set_font(r, size=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("2026")
    set_font(r, size=11)

    doc.add_page_break()
    add_heading(doc, "Apresentação", 1)
    add_body(doc, "Este documento reúne as entregas da Sprint 01 e da Sprint 02 do projeto RotaCerta. A primeira parte registra o problema, os objetivos, os requisitos e o planejamento inicial. A segunda parte apresenta a arquitetura, os modelos de software e dados, os protótipos, a estrutura inicial do banco de dados e a organização do repositório que servirão de base para a implementação.")
    add_body(doc, "A revisão incorpora as orientações recebidas após a Sprint 01. O cronograma passa a adotar cadência semanal, com exceção do período estendido da Sprint 02 informado pela professora. A comparação entre as versões sequencial e paralela do motor de otimização passa a ter métricas, protocolo de medição e uma primeira implementação executável com testes automatizados.")

    add_heading(doc, "Sumário", 1)
    toc_items = [
        "Parte I  Sprint 01 Planejamento do projeto",
        "1  Identificação da equipe",
        "2  Tema, problema, objetivos e público alvo",
        "3  Requisitos funcionais e não funcionais",
        "4  Casos de uso e Product Backlog",
        "5  Cronograma semanal revisado",
        "Parte II  Sprint 02 Arquitetura e modelagem",
        "6  Arquitetura do sistema",
        "7  Diagrama de classes",
        "8  Modelo Entidade Relacionamento",
        "9  Modelo relacional",
        "10  Protótipos das telas",
        "11  Banco de dados e projeto no GitHub",
        "12  Avaliação sequencial e paralela",
        "13  Aderência às orientações da disciplina",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r)

    # Sprint 01
    doc.add_page_break()
    add_heading(doc, "Parte I  Sprint 01 Planejamento do projeto", 1)
    add_heading(doc, "1  Identificação da equipe", 2)
    add_table(doc, ["Integrante", "Matrícula", "Responsabilidade inicial"], [
        ["Álvaro Jordão", "01748200", "Scrum Master e Product Owner; organização, requisitos e priorização"],
        ["Arthur Sales", "01593811", "Backend e núcleo de otimização e paralelização"],
        ["Vinícius Trigueiro", "01794959", "Frontend, banco de dados e documentação"],
    ], [4.2, 2.8, 10.2], 9)
    add_body(doc, "Projeto RotaCerta. Turma 8NA. As responsabilidades organizam o trabalho, mas todos os integrantes participam das decisões, implementação, testes e documentação.")

    add_heading(doc, "2  Tema, problema, objetivos e público alvo", 2)
    add_body(doc, "Tema. Logística e otimização de processos para pequenos negócios com operação própria de entrega.", "Tema.")
    add_body(doc, "Problema. Farmácias, mercados, restaurantes e pequenas distribuidoras frequentemente planejam entregas de forma manual. Essa prática aumenta a distância percorrida, o consumo de combustível e a possibilidade de atraso quando o volume de pedidos cresce.", "Problema.")
    add_body(doc, "Objetivo geral. Desenvolver um sistema web que distribua pedidos entre entregadores e determine uma sequência de paradas que reduza a distância e o tempo total das rotas.", "Objetivo geral.")
    add_body(doc, "Resultados esperados. Centralizar pedidos e entregadores, reduzir distância e tempo em relação à alocação manual, calcular rotas para vários entregadores e disponibilizar uma interface responsiva para a operação diária.", "Resultados esperados.")
    add_body(doc, "Público alvo. Administradores, atendentes e entregadores de pequenos negócios locais. Os clientes finais são beneficiados por entregas mais previsíveis.", "Público alvo.")

    add_heading(doc, "3  Requisitos funcionais e não funcionais", 2)
    add_heading(doc, "Requisitos funcionais", 3)
    functional = [
        ["RF01", "Cadastrar e autenticar usuários com perfis de administrador, atendente e entregador."],
        ["RF02", "Cadastrar pedidos com cliente, endereço, prioridade e horário desejado."],
        ["RF03", "Cadastrar entregadores com disponibilidade e capacidade de carga."],
        ["RF04", "Gerar automaticamente uma rota otimizada para cada entregador."],
        ["RF05", "Atribuir pedidos a entregadores de forma manual ou automática."],
        ["RF06", "Exibir a rota como lista ordenada de paradas e mapa."],
        ["RF07", "Atualizar o status do pedido entre pendente, atribuído, em rota e entregue."],
        ["RF08", "Manter o histórico de entregas realizadas."],
        ["RF09", "Gerar relatório com tempo, distância e comparação da otimização."],
        ["RF10", "Disponibilizar painel administrativo da operação do dia."],
    ]
    add_table(doc, ["ID", "Descrição"], functional, [2.0, 15.2], 9)
    add_heading(doc, "Requisitos não funcionais", 3)
    non_functional = [
        ["RNF01", "Desempenho", "Executar o cálculo em tempo aceitável para vários entregadores e permitir processamento paralelo."],
        ["RNF02", "Segurança", "Autenticar usuários, autorizar por perfil e armazenar senhas com hash seguro."],
        ["RNF03", "Usabilidade", "Manter a interface simples para usuários com diferentes níveis de letramento técnico."],
        ["RNF04", "Portabilidade", "Funcionar em navegadores e adaptar o layout para celulares."],
        ["RNF05", "Escalabilidade", "Suportar o crescimento de pedidos e entregadores sem degradação significativa."],
        ["RNF06", "Confiabilidade", "Preservar pedidos, rotas e histórico com integridade referencial."],
        ["RNF07", "Manutenibilidade", "Manter código organizado, documentado, testável e versionado."],
    ]
    add_table(doc, ["ID", "Categoria", "Descrição"], non_functional, [1.8, 3.2, 12.2], 8.5)

    add_heading(doc, "4  Casos de uso e Product Backlog", 2)
    use_cases = [
        ["UC01", "Atendente ou administrador", "Cadastrar pedido"],
        ["UC02", "Administrador", "Cadastrar entregador"],
        ["UC03", "Administrador", "Gerar rotas para os pedidos do dia"],
        ["UC04", "Entregador", "Visualizar a rota atribuída"],
        ["UC05", "Entregador", "Atualizar o status da entrega"],
        ["UC06", "Administrador", "Consultar métricas das rotas"],
        ["UC07", "Todos os perfis", "Autenticar-se no sistema"],
    ]
    add_table(doc, ["ID", "Ator", "Caso de uso"], use_cases, [2.0, 5.5, 9.7], 9)
    add_heading(doc, "Product Backlog inicial", 3)
    add_body(doc, "Alta prioridade. Autenticação e perfis, CRUD de pedidos, CRUD de entregadores, banco de dados e motor sequencial com vizinho mais próximo e 2-opt.", "Alta prioridade.")
    add_body(doc, "Média prioridade. Visualização das rotas, atualização de status e paralelização para múltiplos entregadores.", "Média prioridade.")
    add_body(doc, "Evolução. Relatórios, painel consolidado, notificações e avaliação futura de C++ com OpenMP ou CUDA.", "Evolução.")

    doc.add_page_break()
    add_heading(doc, "5  Cronograma semanal revisado", 2)
    add_body(doc, "O cronograma abaixo substitui a periodicidade quinzenal apresentada inicialmente. A Sprint 02 mantém o prazo excepcional de 19 de setembro comunicado pela professora; as etapas seguintes usam ciclos semanais.")
    schedule = [
        ["1", "03/09 a 05/09", "Planejamento, requisitos e backlog"],
        ["2", "06/09 a 19/09", "Arquitetura, modelos, protótipos, banco e estrutura do projeto"],
        ["3", "20/09 a 26/09", "Ambiente executável e contrato inicial da API"],
        ["4", "27/09 a 03/10", "Autenticação e autorização por perfil"],
        ["5", "04/10 a 10/10", "CRUD de clientes e pedidos"],
        ["6", "11/10 a 17/10", "CRUD de entregadores e disponibilidade"],
        ["7", "18/10 a 24/10", "Atribuição de pedidos e visualização de rotas"],
        ["8", "25/10 a 31/10", "Vizinho mais próximo em modo sequencial"],
        ["9", "01/11 a 07/11", "2-opt e linha de base de desempenho"],
        ["10", "08/11 a 14/11", "Paralelização para múltiplos entregadores"],
        ["11", "15/11 a 21/11", "Experimentos sequencial versus paralelo"],
        ["12", "22/11 a 28/11", "Relatórios, testes e documentação"],
        ["Final", "29/11 a 05/12", "Correções, vídeos e preparação para a banca"],
    ]
    add_table(doc, ["Sprint", "Período", "Entrega principal"], schedule, [2.0, 3.5, 11.7], 8.5)

    # Sprint 02
    doc.add_page_break()
    add_heading(doc, "Parte II  Sprint 02 Arquitetura e modelagem", 1)
    add_body(doc, "A Sprint 02 transforma os requisitos em uma base técnica verificável. A solução adota uma arquitetura web modular, um modelo de dados relacional e um componente de otimização isolado, permitindo implementar e medir versões sequencial e paralela sem duplicar regras de negócio.")

    add_heading(doc, "6  Arquitetura do sistema", 2)
    tech = [
        ["Interface", "React, TypeScript e Vite", "Aplicação responsiva para administração e operação móvel"],
        ["API", "FastAPI e Python 3.12", "Endpoints REST, validação, OpenAPI e integração com o otimizador"],
        ["Persistência", "PostgreSQL 16", "Integridade, índices, histórico e métricas de execução"],
        ["Mapa", "Leaflet, OpenStreetMap e adaptador OSRM", "Exibição de rotas e matriz de distâncias sem prender o domínio a um provedor"],
        ["Otimização", "Python, 2-opt e processos independentes", "Linha de base sequencial e cálculo paralelo por entregador"],
        ["Ambiente", "Docker Compose", "Inicialização reproduzível de frontend, API e banco"],
    ]
    add_table(doc, ["Camada", "Tecnologia", "Responsabilidade"], tech, [3.0, 4.5, 9.7], 8.5)
    add_figure(doc, "arquitetura.png", "Figura 1  Arquitetura lógica e comunicação entre os componentes")
    add_body(doc, "O navegador envia requisições HTTPS em JSON para a API. A API executa regras de negócio, persiste dados no PostgreSQL e aciona o motor de otimização. O motor recebe pedidos, entregadores e uma matriz de distâncias por meio de uma interface estável. A mesma interface atende aos modos sequencial e paralelo, o que permite comparar desempenho sem alterar os dados de entrada.")
    add_body(doc, "A escolha de Python não pressupõe o uso de inteligência artificial. As orientações da disciplina aceitam inteligência artificial e/ou otimização de processamento, paralelização e alto desempenho. No RotaCerta, o componente avançado é a otimização de rotas integrada ao fluxo principal. React e TypeScript permanecem restritos à interface, enquanto Python executa o núcleo computacional.")

    add_heading(doc, "7  Diagrama de classes", 2)
    add_figure(doc, "diagrama_classes.png", "Figura 2  Classes de domínio e serviço previstas")
    add_body(doc, "Estabelecimento delimita os dados de cada negócio. Usuario representa autenticação e perfil; Entregador complementa um usuário operacional com capacidade e disponibilidade. Pedido registra destino, prioridade, janela de horário e status. Rota agrega paradas ordenadas. OtimizadorRotas executa o mesmo algoritmo nos modos sequencial e paralelo e registra cada ExecucaoOtimizacao.")

    add_heading(doc, "8  Modelo Entidade Relacionamento", 2)
    add_figure(doc, "modelo_entidade_relacionamento.png", "Figura 3  Modelo conceitual com chaves e cardinalidades")
    add_body(doc, "O modelo mantém o estabelecimento como limite organizacional. Um usuário pode possuir um cadastro de entregador. Cada pedido pertence a um cliente e pode ser atribuído a um entregador. Uma rota contém paradas ordenadas; cada parada referencia um pedido. As execuções de otimização registram quantidade de pedidos, entregadores, workers, tempo e distância total.")

    add_heading(doc, "9  Modelo relacional", 2)
    relational = [
        ["establishments", "id", "-", "Negócio, endereço e coordenadas do depósito"],
        ["users", "id", "establishment_id", "Credenciais, perfil e situação do usuário"],
        ["couriers", "id", "establishment_id, user_id", "Capacidade e disponibilidade do entregador"],
        ["customers", "id", "establishment_id", "Cliente associado ao estabelecimento"],
        ["orders", "id", "establishment_id, customer_id, assigned_courier_id", "Destino, peso, prioridade, janela e status"],
        ["routes", "id", "establishment_id, courier_id", "Rota diária, algoritmo, modo e totais"],
        ["route_stops", "id", "route_id, order_id", "Sequência e acompanhamento de cada parada"],
        ["optimization_runs", "id", "establishment_id", "Métricas de execuções sequenciais e paralelas"],
    ]
    add_table(doc, ["Tabela", "PK", "FK", "Finalidade"], relational, [3.3, 1.6, 5.2, 7.1], 8)
    add_body(doc, "As restrições CHECK limitam perfis, estados e modos de execução. As chaves estrangeiras preservam a integridade entre usuários, pedidos, entregadores e rotas. Índices cobrem as consultas mais frequentes por estabelecimento, status, entregador, data e modo de execução.")

    add_heading(doc, "10  Protótipos das telas principais", 2)
    add_body(doc, "Os wireframes priorizam navegação direta e leitura em dispositivos móveis. A primeira sequência cobre autenticação, painel diário e cadastro de pedido. A segunda cobre geração de rotas, consulta da rota pelo entregador e atualização de status.")
    add_figure(doc, "prototipos_administracao.png", "Figura 4  Acesso, painel administrativo e cadastro de pedido")
    add_figure(doc, "prototipos_rotas.png", "Figura 5  Planejamento de rotas e fluxo móvel do entregador")
    prototype_map = [
        ["Acesso", "UC07", "Login e identificação do perfil"],
        ["Painel do dia", "RF10", "Pedidos, situação operacional e atalhos"],
        ["Novo pedido", "UC01 / RF02", "Cliente, destino, carga, prioridade e janela"],
        ["Planejamento", "UC03 / RF04", "Mapa, totais e comando de otimização"],
        ["Minha rota", "UC04 / RF06", "Sequência, mapa e próxima parada"],
        ["Atualizar entrega", "UC05 / RF07", "Transição de status e observação"],
    ]
    add_table(doc, ["Tela", "Requisito", "Objetivo"], prototype_map, [4.0, 3.2, 10.0], 8.5)

    doc.add_page_break()
    add_heading(doc, "11  Banco de dados e projeto no GitHub", 2)
    add_heading(doc, "Banco de dados criado", 3)
    add_body(doc, "A estrutura inicial está implementada no arquivo database/schema.sql para PostgreSQL 16. O serviço db do arquivo compose.yaml executa o esquema automaticamente na primeira inicialização do volume. O modelo inclui oito tabelas, restrições de integridade e cinco índices iniciais.")
    add_numbered(doc, [
        "Copiar .env.example para .env.",
        "Executar docker compose up --build.",
        "Aguardar o health check do PostgreSQL.",
        "Consultar a API em http://localhost:8000/health e a documentação em http://localhost:8000/docs.",
    ])
    add_heading(doc, "Estrutura do repositório", 3)
    repository_rows = [
        ["backend/", "API FastAPI e limite do módulo de otimização"],
        ["backend/app/optimizer/", "Vizinho mais próximo, 2-opt e execução sequencial ou paralela"],
        ["backend/tests/", "Testes automatizados do núcleo de roteirização"],
        ["frontend/", "Aplicação React e TypeScript"],
        ["database/", "DDL, instruções e inicialização do PostgreSQL"],
        ["docs/sprint-02/assets/", "Diagramas e protótipos incorporados ao documento"],
        ["tools/documentation/", "Geração reproduzível dos materiais técnicos"],
        ["compose.yaml", "Orquestração local dos três serviços"],
        ["CONTRIBUTING.md", "Padrões de branch, commit e revisão"],
    ]
    add_table(doc, ["Caminho", "Conteúdo"], repository_rows, [6.0, 11.2], 9)
    add_body(doc, "Repositório oficial. https://github.com/alvccpj/rota-certa", "Repositório oficial.")
    branch_rows = [
        ["master", "Versão estável e entregas aprovadas"],
        ["sprint/02-arquitetura-modelagem", "Integração da Sprint 02"],
        ["dev/alvaro", "Ambiente pessoal de Álvaro"],
        ["dev/arthur", "Ambiente pessoal de Arthur"],
        ["dev/vinicius", "Ambiente pessoal de Vinícius"],
    ]
    add_table(doc, ["Branch", "Finalidade"], branch_rows, [7.0, 10.2], 9)
    add_body(doc, "O fluxo recomendado usa branches curtas por tarefa, pull request e revisão de outro integrante. As branches pessoais identificam ambientes acadêmicos, mas não substituem branches de feature, correção ou documentação.")

    doc.add_page_break()
    add_heading(doc, "12  Avaliação sequencial e paralela", 2)
    add_body(doc, "A orientação da Sprint 01 exige manter a comparação entre as versões sequencial e paralela. A primeira implementação já executa vizinho mais próximo e refinamento 2-opt nos dois modos. As versões usam os mesmos pedidos, entregadores e coordenadas; a única variável controlada é o modo de execução.")
    metrics = [
        ["Tempo de execução", "Milissegundos", "Média, mediana e percentil 95 de dez repetições"],
        ["Distância total", "Quilômetros", "Confirma que a paralelização não altera a qualidade da rota"],
        ["Speedup", "Tseq / Tpar", "Ganho obtido com a execução paralela"],
        ["Eficiência", "Speedup / workers", "Uso relativo dos processos disponíveis"],
        ["Escala", "Pedidos e entregadores", "Cenários de 50, 100, 250 e 500 pedidos com 2, 4 e 8 entregadores"],
    ]
    add_table(doc, ["Métrica", "Cálculo ou unidade", "Critério"], metrics, [4.0, 4.0, 9.2], 8.5)
    add_body(doc, "Cada cenário terá uma execução de aquecimento e dez repetições válidas. O relatório registrará processador, memória, sistema operacional, versão do código, número de workers e semente dos dados. A tabela optimization_runs armazenará cada medição para permitir auditoria e reprodução.")

    add_heading(doc, "Evidência prática", 3)
    add_body(doc, "O arquivo backend/app/optimizer/routing.py contém a heurística determinística, o cálculo de distância geodésica, o refinamento 2-opt e a execução com processos independentes. O endpoint POST /optimizer/compare retorna rotas, distâncias, tempos, workers, speedup e a confirmação de equivalência entre os modos. Cinco testes automatizados validam rota vazia, preservação das paradas, melhoria do 2-opt, resultado da otimização e equivalência sequencial/paralela.")

    doc.add_page_break()
    add_heading(doc, "13  Aderência às orientações da disciplina", 2)
    alignment = [
        ["Projeto integrado", "Atendido", "Um único sistema reúne interface, API, banco e otimização"],
        ["Componente avançado", "Atendido", "Otimização e paralelização fazem parte da geração de rotas"],
        ["Tecnologia prioritária", "Atendido", "Python no núcleo; JavaScript e TypeScript apenas na interface"],
        ["Integração não superficial", "Atendido", "O endpoint compara os modos sobre os mesmos dados"],
        ["Evolução prática", "Atendido", "Código executável e cinco testes automatizados no repositório"],
        ["GitHub atualizado", "Atendido", "Branch da Sprint 02 e Pull Request número 1"],
        ["Participação da equipe", "Contínuo", "Cada integrante deve registrar contribuições próprias nas próximas tarefas"],
    ]
    add_table(doc, ["Orientação", "Situação", "Evidência"], alignment, [4.1, 2.8, 10.3], 8.5)
    add_body(doc, "Os requisitos de software completo, autenticação, perfis, funcionalidades integradas e demonstração final pertencem à evolução do semestre. A Sprint 02 estabelece a base verificável para essas implementações sem declarar como concluído o produto final.")

    add_heading(doc, "Situação da Sprint 02", 2)
    checklist = [
        ["Arquitetura do sistema", "Concluído", "Tecnologias, responsabilidades, comunicação e componente avançado definidos"],
        ["Diagrama de classes", "Concluído", "Classes, atributos, métodos e relacionamentos apresentados"],
        ["MER", "Concluído", "Entidades, chaves e cardinalidades apresentadas"],
        ["Modelo relacional", "Concluído", "Tabelas, PKs, FKs e finalidade documentadas"],
        ["Protótipos", "Concluído", "Seis telas principais em baixa fidelidade"],
        ["Banco de dados", "Concluído", "DDL PostgreSQL e inicialização via Compose"],
        ["Motor de otimização", "Concluído", "Primeira versão sequencial/paralela com cinco testes"],
        ["GitHub", "Concluído", "Estrutura, padrões, branch da Sprint e Pull Request número 1"],
        ["Identificação do grupo", "Concluído", "Grupo 04 informado na capa e no nome do arquivo"],
    ]
    add_table(doc, ["Entrega", "Situação", "Evidência"], checklist, [4.0, 2.8, 10.4], 8.5)

    add_heading(doc, "Próxima etapa", 2)
    add_body(doc, "A Sprint 03 deve validar o ambiente em uma máquina da equipe, confirmar a criação do PostgreSQL e integrar o frontend aos endpoints de saúde e comparação do otimizador. A equipe também deve transformar os casos de uso prioritários em tarefas semanais com critérios de aceite e registrar contribuições individuais no GitHub.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

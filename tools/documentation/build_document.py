"""Gera o documento técnico cumulativo das Sprints 01 a 05.

O mesmo conteúdo é renderizado em HTML, PDF (Chrome headless via Playwright) e
DOCX (python-docx). As evidências de cada Sprint são produzidas antes pelos
scripts capture_sprintNN_evidence.py com o sistema em execução.

    python tools/documentation/build_document.py
"""

from dataclasses import dataclass, field
from html import escape
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from build_sprint02_document import (
    GRAY,
    add_body,
    add_bullets,
    add_heading,
    add_numbered,
    add_table,
    set_cell_margins,
    set_cell_shading,
    set_font,
    set_table_borders,
    setup_document,
)

ROOT = Path(__file__).resolve().parents[2]
SPRINT02 = ROOT / "docs" / "sprint-02" / "assets"
SPRINT03 = ROOT / "docs" / "sprint-03" / "assets"
EVIDENCE = ROOT / "docs" / "sprint-03" / "evidencias"
SPRINT04 = ROOT / "docs" / "sprint-04" / "assets"
EVIDENCE04 = ROOT / "docs" / "sprint-04" / "evidencias"
SPRINT05 = ROOT / "docs" / "sprint-05" / "assets"
EVIDENCE05 = ROOT / "docs" / "sprint-05" / "evidencias"
NAME = "Grupo_04_RotaCerta_Sprints_01_a_05"
TITLE = "Documento Técnico Cumulativo das Sprints 01 a 05"
FOOTER = "RotaCerta  |  Grupo 04  |  Sprints 01 a 05  |  Página "
REPOSITORY = "https://github.com/alvccpj/rota-certa"
TEAM = (
    ("Álvaro Jordão", "01748200"),
    ("Arthur Sales", "01593811"),
    ("Vinícius Trigueiro", "01794959"),
    ("William Coelho de Morais", "01263977"),
)


# Blocos de conteúdo

@dataclass
class H:
    text: str
    level: int = 2


@dataclass
class P:
    text: str
    lead: str | None = None


@dataclass
class Bullets:
    items: list[str]


@dataclass
class Numbered:
    items: list[str]


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    widths: list[float]
    font_size: float = 8.5


@dataclass
class Figure:
    paths: list[Path]
    caption: str
    width: float = 100  # porcentagem da largura útil, por imagem


@dataclass
class Code:
    text: str
    caption: str


@dataclass
class PageBreak:
    pass


@dataclass
class Document:
    blocks: list = field(default_factory=list)

    def add(self, *blocks) -> None:
        self.blocks.extend(blocks)


def evidence(name: str) -> str:
    return (EVIDENCE / name).read_text(encoding="utf-8").rstrip()


def s2(name: str) -> Path:
    return SPRINT02 / name


def s3(name: str) -> Path:
    return SPRINT03 / name


def s4(name: str) -> Path:
    return SPRINT04 / name


def evidence04(name: str) -> str:
    return (EVIDENCE04 / name).read_text(encoding="utf-8").rstrip()


def s5(name: str) -> Path:
    return SPRINT05 / name


def evidence05(name: str) -> str:
    return (EVIDENCE05 / name).read_text(encoding="utf-8").rstrip()


# Conteúdo

PRESENTATION = [
    "Este documento reúne as entregas das Sprints 01 a 05 do projeto RotaCerta. A Parte I registra o problema, os objetivos, os requisitos e o planejamento inicial. A Parte II apresenta a arquitetura, os modelos de software e de dados, os protótipos e a estrutura inicial do banco. A Parte III demonstra a estrutura inicial funcionando: banco conectado, login, cadastro de usuários, controle de perfis e o CRUD de pedidos. A Parte IV apresenta o primeiro módulo completo, Pedidos e entregas, com persistência, validações, mensagens de erro e navegação entre as telas. A Parte V apresenta o segundo módulo, Roteirização e desempenho, que integra o núcleo de otimização aos pedidos reais, mostra as rotas no mapa e compara as execuções sequencial, paralela em CPU e na GPU com CUDA.",
    "As Partes I a IV mantêm o conteúdo entregue anteriormente. Os ajustes de planejamento, arquitetura e modelagem da Sprint 04 estão na seção 32 e os da Sprint 05 na seção 43. Todas as telas e respostas mostradas nas Partes III a V foram capturadas do sistema em execução.",
]

TOC = [
    "Parte I  Sprint 01 Planejamento do projeto",
    "1  Identificação da equipe",
    "2  Tema, problema, objetivos e público alvo",
    "3  Requisitos funcionais e não funcionais",
    "4  Casos de uso e Product Backlog",
    "5  Cronograma semanal revisado",
    "Parte II  Sprint 02 Arquitetura e modelagem",
    "6  Arquitetura do sistema",
    "7  Diagrama de classes",
    "8  Modelo Entidade Relacionamento",
    "9  Modelo relacional",
    "10  Protótipos das telas",
    "11  Banco de dados e projeto no GitHub",
    "12  Avaliação sequencial e paralela",
    "13  Aderência às orientações da disciplina",
    "Parte III  Sprint 03 Estrutura inicial funcionando",
    "14  Estrutura implementada",
    "15  Banco de dados conectado",
    "16  Login funcional",
    "17  Cadastro de usuários",
    "18  Controle de perfis",
    "19  CRUD principal: pedidos",
    "20  Execução local",
    "21  Testes automatizados",
    "22  Ajustes no planejamento e na modelagem",
    "23  Repositório GitHub",
    "24  Dificuldades encontradas e próximos passos",
    "Parte IV  Sprint 04 Primeiro módulo completo",
    "25  Módulo implementado: Pedidos e entregas",
    "26  Evidências do módulo funcionando",
    "27  Persistência de dados",
    "28  Validações",
    "29  Mensagens de erro",
    "30  Navegação entre as telas",
    "31  Commits organizados e repositório",
    "32  Ajustes no planejamento, na arquitetura e na modelagem",
    "33  Testes automatizados",
    "34  Dificuldades encontradas e próximos passos",
    "Parte V  Sprint 05 Segundo módulo funcionando",
    "35  Módulo implementado: Roteirização e desempenho",
    "36  Evidências das funcionalidades",
    "37  Integração com o banco de dados",
    "38  Regras de negócio",
    "39  Núcleo de otimização e paralelização com CUDA",
    "40  Comparação sequencial, paralela e GPU",
    "41  Testes das funcionalidades",
    "42  Bugs encontrados e correções",
    "43  Ajustes no planejamento, na arquitetura e na modelagem",
    "44  Repositório GitHub",
    "45  Dificuldades encontradas e próximos passos",
]


def sprint01(doc: Document) -> None:
    doc.add(
        PageBreak(),
        H("Parte I  Sprint 01 Planejamento do projeto", 1),
        H("1  Identificação da equipe"),
        Table(
            ["Integrante", "Matrícula", "Responsabilidade inicial"],
            [
                ["Álvaro Jordão", "01748200", "Scrum Master e Product Owner; organização, requisitos e priorização"],
                ["Arthur Sales", "01593811", "Backend e núcleo de otimização e paralelização"],
                ["Vinícius Trigueiro", "01794959", "Frontend, banco de dados e documentação"],
                ["William Coelho de Morais", "01263977", "Desenvolvimento e testes"],
            ],
            [4.2, 2.8, 10.2],
            9,
        ),
        P("Projeto RotaCerta. Turma 8NA. As responsabilidades organizam o trabalho, mas todos os integrantes participam das decisões, implementação, testes e documentação."),
        H("2  Tema, problema, objetivos e público alvo"),
        P("Tema. Logística e otimização de processos para pequenos negócios com operação própria de entrega.", "Tema."),
        P("Problema. Farmácias, mercados, restaurantes e pequenas distribuidoras frequentemente planejam entregas de forma manual. Essa prática aumenta a distância percorrida, o consumo de combustível e a possibilidade de atraso quando o volume de pedidos cresce.", "Problema."),
        P("Objetivo geral. Desenvolver um sistema web que distribua pedidos entre entregadores e determine uma sequência de paradas que reduza a distância e o tempo total das rotas.", "Objetivo geral."),
        P("Resultados esperados. Centralizar pedidos e entregadores, reduzir distância e tempo em relação à alocação manual, calcular rotas para vários entregadores e disponibilizar uma interface responsiva para a operação diária.", "Resultados esperados."),
        P("Público alvo. Administradores, atendentes e entregadores de pequenos negócios locais. Os clientes finais são beneficiados por entregas mais previsíveis.", "Público alvo."),
        H("3  Requisitos funcionais e não funcionais"),
        H("Requisitos funcionais", 3),
        Table(
            ["ID", "Descrição"],
            [
                ["RF01", "Cadastrar e autenticar usuários com perfis de administrador, atendente e entregador."],
                ["RF02", "Cadastrar pedidos com cliente, endereço, prioridade e horário desejado."],
                ["RF03", "Cadastrar entregadores com disponibilidade e capacidade de carga."],
                ["RF04", "Gerar automaticamente uma rota otimizada para cada entregador."],
                ["RF05", "Atribuir pedidos a entregadores de forma manual ou automática."],
                ["RF06", "Exibir a rota como lista ordenada de paradas e mapa."],
                ["RF07", "Atualizar o status do pedido entre pendente, atribuído, em rota e entregue."],
                ["RF08", "Manter o histórico de entregas realizadas."],
                ["RF09", "Gerar relatório com tempo, distância e comparação da otimização."],
                ["RF10", "Disponibilizar painel administrativo da operação do dia."],
            ],
            [2.0, 15.2],
            9,
        ),
        H("Requisitos não funcionais", 3),
        Table(
            ["ID", "Categoria", "Descrição"],
            [
                ["RNF01", "Desempenho", "Executar o cálculo em tempo aceitável para vários entregadores e permitir processamento paralelo."],
                ["RNF02", "Segurança", "Autenticar usuários, autorizar por perfil e armazenar senhas com hash seguro."],
                ["RNF03", "Usabilidade", "Manter a interface simples para usuários com diferentes níveis de letramento técnico."],
                ["RNF04", "Portabilidade", "Funcionar em navegadores e adaptar o layout para celulares."],
                ["RNF05", "Escalabilidade", "Suportar o crescimento de pedidos e entregadores sem degradação significativa."],
                ["RNF06", "Confiabilidade", "Preservar pedidos, rotas e histórico com integridade referencial."],
                ["RNF07", "Manutenibilidade", "Manter código organizado, documentado, testável e versionado."],
            ],
            [1.8, 3.2, 12.2],
        ),
        H("4  Casos de uso e Product Backlog"),
        Table(
            ["ID", "Ator", "Caso de uso"],
            [
                ["UC01", "Atendente ou administrador", "Cadastrar pedido"],
                ["UC02", "Administrador", "Cadastrar entregador"],
                ["UC03", "Administrador", "Gerar rotas para os pedidos do dia"],
                ["UC04", "Entregador", "Visualizar a rota atribuída"],
                ["UC05", "Entregador", "Atualizar o status da entrega"],
                ["UC06", "Administrador", "Consultar métricas das rotas"],
                ["UC07", "Todos os perfis", "Autenticar-se no sistema"],
            ],
            [2.0, 5.5, 9.7],
            9,
        ),
        H("Product Backlog inicial", 3),
        P("Alta prioridade. Autenticação e perfis, CRUD de pedidos, CRUD de entregadores, banco de dados e motor sequencial com vizinho mais próximo e 2-opt.", "Alta prioridade."),
        P("Média prioridade. Visualização das rotas, atualização de status e paralelização para múltiplos entregadores.", "Média prioridade."),
        P("Evolução. Relatórios, painel consolidado, notificações e avaliação futura de C++ com OpenMP ou CUDA.", "Evolução."),
        PageBreak(),
        H("5  Cronograma semanal revisado"),
        P("O cronograma abaixo substituiu, na Sprint 02, a periodicidade quinzenal apresentada inicialmente. A Sprint 02 manteve o prazo excepcional de 19 de setembro comunicado pela professora; as etapas seguintes usam ciclos semanais. A versão vigente, com os ajustes das Sprints 03 e 04, está na seção 32."),
        Table(
            ["Sprint", "Período", "Entrega principal"],
            [
                ["1", "03/09 a 05/09", "Planejamento, requisitos e backlog"],
                ["2", "06/09 a 19/09", "Arquitetura, modelos, protótipos, banco e estrutura do projeto"],
                ["3", "20/09 a 26/09", "Ambiente executável e contrato inicial da API"],
                ["4", "27/09 a 03/10", "Autenticação e autorização por perfil"],
                ["5", "04/10 a 10/10", "CRUD de clientes e pedidos"],
                ["6", "11/10 a 17/10", "CRUD de entregadores e disponibilidade"],
                ["7", "18/10 a 24/10", "Atribuição de pedidos e visualização de rotas"],
                ["8", "25/10 a 31/10", "Vizinho mais próximo em modo sequencial"],
                ["9", "01/11 a 07/11", "2-opt e linha de base de desempenho"],
                ["10", "08/11 a 14/11", "Paralelização para múltiplos entregadores"],
                ["11", "15/11 a 21/11", "Experimentos sequencial versus paralelo"],
                ["12", "22/11 a 28/11", "Relatórios, testes e documentação"],
                ["Final", "29/11 a 05/12", "Correções, vídeos e preparação para a banca"],
            ],
            [2.0, 3.5, 11.7],
        ),
    )


def sprint02(doc: Document) -> None:
    doc.add(
        PageBreak(),
        H("Parte II  Sprint 02 Arquitetura e modelagem", 1),
        P("A Sprint 02 transforma os requisitos em uma base técnica verificável. A solução adota uma arquitetura web modular, um modelo de dados relacional e um componente de otimização isolado, permitindo implementar e medir versões sequencial e paralela sem duplicar regras de negócio."),
        H("6  Arquitetura do sistema"),
        Table(
            ["Camada", "Tecnologia", "Responsabilidade"],
            [
                ["Interface", "React, TypeScript e Vite", "Aplicação responsiva para administração e operação móvel"],
                ["API", "FastAPI e Python 3.12", "Endpoints REST, validação, OpenAPI e integração com o otimizador"],
                ["Persistência", "PostgreSQL 16", "Integridade, índices, histórico e métricas de execução"],
                ["Mapa", "Leaflet, OpenStreetMap e adaptador OSRM", "Exibição de rotas e matriz de distâncias sem prender o domínio a um provedor"],
                ["Otimização", "Python, 2-opt e processos independentes", "Linha de base sequencial e cálculo paralelo por entregador"],
                ["Ambiente", "Docker Compose", "Inicialização reproduzível de frontend, API e banco"],
            ],
            [3.0, 4.5, 9.7],
        ),
        Figure([s2("arquitetura.png")], "Arquitetura lógica e comunicação entre os componentes"),
        P("O navegador envia requisições HTTPS em JSON para a API. A API executa regras de negócio, persiste dados no PostgreSQL e aciona o motor de otimização. O motor recebe pedidos, entregadores e uma matriz de distâncias por meio de uma interface estável. A mesma interface atende aos modos sequencial e paralelo, o que permite comparar desempenho sem alterar os dados de entrada."),
        P("A escolha de Python não pressupõe o uso de inteligência artificial. As orientações da disciplina aceitam inteligência artificial e/ou otimização de processamento, paralelização e alto desempenho. No RotaCerta, o componente avançado é a otimização de rotas integrada ao fluxo principal. React e TypeScript permanecem restritos à interface, enquanto Python executa o núcleo computacional."),
        H("7  Diagrama de classes"),
        Figure([s2("diagrama_classes.png")], "Classes de domínio e serviço previstas"),
        P("Estabelecimento delimita os dados de cada negócio. Usuario representa autenticação e perfil; Entregador complementa um usuário operacional com capacidade e disponibilidade. Pedido registra destino, prioridade, janela de horário e status. Rota agrega paradas ordenadas. OtimizadorRotas executa o mesmo algoritmo nos modos sequencial e paralelo e registra cada ExecucaoOtimizacao."),
        H("8  Modelo Entidade Relacionamento"),
        Figure([s2("modelo_entidade_relacionamento.png")], "Modelo conceitual com chaves e cardinalidades"),
        P("O modelo mantém o estabelecimento como limite organizacional. Um usuário pode possuir um cadastro de entregador. Cada pedido pertence a um cliente e pode ser atribuído a um entregador. Uma rota contém paradas ordenadas; cada parada referencia um pedido. As execuções de otimização registram quantidade de pedidos, entregadores, workers, tempo e distância total."),
        H("9  Modelo relacional"),
        Table(
            ["Tabela", "PK", "FK", "Finalidade"],
            [
                ["establishments", "id", "-", "Negócio, endereço e coordenadas do depósito"],
                ["users", "id", "establishment_id", "Credenciais, perfil e situação do usuário"],
                ["couriers", "id", "establishment_id, user_id", "Capacidade e disponibilidade do entregador"],
                ["customers", "id", "establishment_id", "Cliente associado ao estabelecimento"],
                ["orders", "id", "establishment_id, customer_id, assigned_courier_id", "Destino, peso, prioridade, janela e status"],
                ["routes", "id", "establishment_id, courier_id", "Rota diária, algoritmo, modo e totais"],
                ["route_stops", "id", "route_id, order_id", "Sequência e acompanhamento de cada parada"],
                ["optimization_runs", "id", "establishment_id", "Métricas de execuções sequenciais e paralelas"],
            ],
            [3.3, 1.6, 5.2, 7.1],
            8,
        ),
        P("As restrições CHECK limitam perfis, estados e modos de execução. As chaves estrangeiras preservam a integridade entre usuários, pedidos, entregadores e rotas. Índices cobrem as consultas mais frequentes por estabelecimento, status, entregador, data e modo de execução."),
        H("10  Protótipos das telas principais"),
        P("Os wireframes priorizam navegação direta e leitura em dispositivos móveis. A primeira sequência cobre autenticação, painel diário e cadastro de pedido. A segunda cobre geração de rotas, consulta da rota pelo entregador e atualização de status."),
        Figure([s2("prototipos_administracao.png")], "Acesso, painel administrativo e cadastro de pedido"),
        Figure([s2("prototipos_rotas.png")], "Planejamento de rotas e fluxo móvel do entregador"),
        Table(
            ["Tela", "Requisito", "Objetivo"],
            [
                ["Acesso", "UC07", "Login e identificação do perfil"],
                ["Painel do dia", "RF10", "Pedidos, situação operacional e atalhos"],
                ["Novo pedido", "UC01 / RF02", "Cliente, destino, carga, prioridade e janela"],
                ["Planejamento", "UC03 / RF04", "Mapa, totais e comando de otimização"],
                ["Minha rota", "UC04 / RF06", "Sequência, mapa e próxima parada"],
                ["Atualizar entrega", "UC05 / RF07", "Transição de status e observação"],
            ],
            [4.0, 3.2, 10.0],
        ),
        PageBreak(),
        H("11  Banco de dados e projeto no GitHub"),
        H("Banco de dados criado", 3),
        P("A estrutura inicial está implementada no arquivo database/schema.sql para PostgreSQL 16. O serviço db do arquivo compose.yaml executa o esquema automaticamente na primeira inicialização do volume. O modelo inclui oito tabelas, restrições de integridade e cinco índices iniciais."),
        H("Estrutura do repositório na Sprint 02", 3),
        Table(
            ["Caminho", "Conteúdo"],
            [
                ["backend/", "API FastAPI e limite do módulo de otimização"],
                ["backend/app/optimizer/", "Vizinho mais próximo, 2-opt e execução sequencial ou paralela"],
                ["backend/tests/", "Testes automatizados do núcleo de roteirização"],
                ["frontend/", "Aplicação React e TypeScript"],
                ["database/", "DDL, instruções e inicialização do PostgreSQL"],
                ["docs/sprint-02/assets/", "Diagramas e protótipos incorporados ao documento"],
                ["tools/documentation/", "Geração reproduzível dos materiais técnicos"],
                ["compose.yaml", "Orquestração local dos três serviços"],
                ["CONTRIBUTING.md", "Padrões de branch, commit e revisão"],
            ],
            [6.0, 11.2],
            9,
        ),
        P("O fluxo recomendado usa branches curtas por tarefa, pull request e revisão de outro integrante. As branches pessoais dev/alvaro, dev/arthur e dev/vinicius identificam ambientes acadêmicos, mas não substituem branches de feature, correção ou documentação."),
        H("12  Avaliação sequencial e paralela"),
        P("A orientação da Sprint 01 exige manter a comparação entre as versões sequencial e paralela. A primeira implementação já executa vizinho mais próximo e refinamento 2-opt nos dois modos. As versões usam os mesmos pedidos, entregadores e coordenadas; a única variável controlada é o modo de execução."),
        Table(
            ["Métrica", "Cálculo ou unidade", "Critério"],
            [
                ["Tempo de execução", "Milissegundos", "Média, mediana e percentil 95 de dez repetições"],
                ["Distância total", "Quilômetros", "Confirma que a paralelização não altera a qualidade da rota"],
                ["Speedup", "Tseq / Tpar", "Ganho obtido com a execução paralela"],
                ["Eficiência", "Speedup / workers", "Uso relativo dos processos disponíveis"],
                ["Escala", "Pedidos e entregadores", "Cenários de 50, 100, 250 e 500 pedidos com 2, 4 e 8 entregadores"],
            ],
            [4.0, 4.0, 9.2],
        ),
        P("Cada cenário terá uma execução de aquecimento e dez repetições válidas. O relatório registrará processador, memória, sistema operacional, versão do código, número de workers e semente dos dados. A tabela optimization_runs armazenará cada medição para permitir auditoria e reprodução."),
        P("Evidência prática. O arquivo backend/app/optimizer/routing.py contém a heurística determinística, o cálculo de distância geodésica, o refinamento 2-opt e a execução com processos independentes. O endpoint POST /optimizer/compare retorna rotas, distâncias, tempos, workers, speedup e a confirmação de equivalência entre os modos. Cinco testes automatizados validam rota vazia, preservação das paradas, melhoria do 2-opt, resultado da otimização e equivalência sequencial/paralela.", "Evidência prática."),
        H("13  Aderência às orientações da disciplina"),
        Table(
            ["Orientação", "Situação na Sprint 02", "Evidência"],
            [
                ["Projeto integrado", "Atendido", "Um único sistema reúne interface, API, banco e otimização"],
                ["Componente avançado", "Atendido", "Otimização e paralelização fazem parte da geração de rotas"],
                ["Tecnologia prioritária", "Atendido", "Python no núcleo; JavaScript e TypeScript apenas na interface"],
                ["Integração não superficial", "Atendido", "O endpoint compara os modos sobre os mesmos dados"],
                ["Evolução prática", "Atendido", "Código executável e cinco testes automatizados no repositório"],
                ["GitHub atualizado", "Atendido", "Branch da Sprint 02 e Pull Request número 1"],
                ["Participação da equipe", "Contínuo", "Cada integrante deve registrar contribuições próprias nas próximas tarefas"],
            ],
            [4.1, 3.2, 9.9],
        ),
    )


def sprint03(doc: Document) -> None:
    doc.add(
        PageBreak(),
        H("Parte III  Sprint 03 Estrutura inicial funcionando", 1),
        P("A Sprint 03 transforma a arquitetura e o modelo de dados da Sprint 02 em um sistema executável. Frontend, API e PostgreSQL passaram a funcionar integrados: o usuário faz login, recebe permissões de acordo com o seu perfil e cadastra, consulta, edita e exclui pedidos gravados no banco. Todas as telas e respostas mostradas nesta parte foram capturadas com o sistema em execução local."),
        H("14  Estrutura implementada"),
        P("A aplicação mantém as três camadas definidas na Sprint 02. O frontend React consome a API REST com requisições JSON autenticadas por token. A API valida os dados, aplica as regras de cada perfil e grava as informações no PostgreSQL por meio do SQLAlchemy."),
        Table(
            ["Componente", "Arquivo", "Responsabilidade"],
            [
                ["Configuração", "backend/app/config.py", "Lê a URL do banco, a chave do token, as origens permitidas e as opções de inicialização"],
                ["Conexão", "backend/app/database.py", "Cria o pool de conexões, verifica o banco e aplica o esquema quando necessário"],
                ["Modelos", "backend/app/models.py", "Mapeia as tabelas do schema.sql para classes Python"],
                ["Segurança", "backend/app/security.py", "Hash Argon2, emissão e validação do JWT e verificação de perfil"],
                ["Autenticação", "backend/app/routers/auth.py", "Login, cadastro de estabelecimento e dados do usuário logado"],
                ["Usuários", "backend/app/routers/users.py", "Cadastro, edição e desativação de usuários; lista de entregadores"],
                ["Pedidos", "backend/app/routers/orders.py", "CRUD de pedidos e mudança de situação"],
                ["Dados iniciais", "backend/app/seed.py", "Estabelecimento, usuários e pedidos de demonstração"],
                ["Interface", "frontend/src/pages/", "Login, pedidos, entregas do entregador e usuários"],
                ["Testes", "backend/tests/", "16 testes da API e 5 do otimizador"],
            ],
            [3.0, 5.2, 9.0],
        ),
        H("Fluxo de uma requisição autenticada", 3),
        Numbered([
            "O usuário informa e-mail e senha na tela de login.",
            "A API compara a senha com o hash Argon2 gravado em users.password_hash e devolve um token JWT válido por oito horas.",
            "O frontend envia o token no cabeçalho Authorization de cada requisição.",
            "A API valida o token, carrega o usuário do banco e confere se o perfil tem permissão para a rota.",
            "As consultas são sempre filtradas pelo estabelecimento do usuário; entregadores recebem apenas os pedidos atribuídos a eles.",
        ]),
        H("15  Banco de dados conectado"),
        P("A API se conecta ao PostgreSQL pela URL da variável DATABASE_URL, usando o driver psycopg 3 e o SQLAlchemy 2. O pool de conexões testa a conexão antes de cada uso. Na execução com Docker, o PostgreSQL aplica database/schema.sql na primeira inicialização; na execução sem Docker, a própria API aplica o mesmo arquivo quando encontra o banco vazio. Nos dois casos a estrutura é exatamente a modelada na Sprint 02."),
        P("O endpoint GET /health executa uma consulta no banco e informa o tipo e a versão do servidor. Se o banco estiver indisponível, a resposta passa a ser HTTP 503 com database igual a unavailable."),
        Figure([s3("s03_01_health.png")], "Endpoint de verificação confirmando a conexão da API com o PostgreSQL", 75),
        Code(evidence("sql_tabelas.txt"), "Consulta às oito tabelas do banco rotacerta com a quantidade de registros"),
        Figure([s3("s03_02_swagger.png")], "Documentação interativa da API gerada pelo FastAPI, com as rotas de autenticação, usuários e pedidos"),
        H("16  Login funcional"),
        P("A tela de login autentica os usuários cadastrados no banco. O e-mail é comparado sem diferenciar maiúsculas de minúsculas e a senha é verificada contra o hash Argon2; a senha em texto nunca é gravada. Credenciais incorretas retornam HTTP 401 com a mensagem E-mail ou senha inválidos. Usuários desativados recebem HTTP 403. A sessão fica salva no navegador e é validada novamente na API sempre que a página é aberta."),
        Figure([s3("s03_03_login.png")], "Tela de login com atalhos para as contas de demonstração"),
        Figure([s3("s03_04_login_erro.png")], "Tentativa de login com senha incorreta"),
        Code(evidence("sql_usuarios.txt"), "Usuários gravados no banco: apenas o hash da senha é armazenado"),
        H("17  Cadastro de usuários"),
        P("O cadastro acontece de duas formas. Um novo negócio cria a própria conta pela opção Cadastrar meu negócio, que grava o estabelecimento e o seu primeiro administrador. A partir daí, o administrador cadastra atendentes e entregadores na tela Usuários. Para entregadores, o sistema exige a capacidade de carga e cria também o registro na tabela couriers."),
        Bullets([
            "O e-mail é único em todo o sistema; uma duplicidade retorna HTTP 409.",
            "A senha tem no mínimo oito caracteres e é gravada com hash Argon2.",
            "Usuários não são apagados: a desativação bloqueia o login e preserva o histórico de pedidos.",
            "O administrador não pode desativar a si mesmo nem remover o próprio perfil de administrador.",
            "Cada estabelecimento enxerga apenas os próprios usuários e pedidos.",
        ]),
        Figure([s3("s03_05_cadastro_negocio.png")], "Cadastro de um novo estabelecimento com o seu administrador"),
        Figure([s3("s03_06_negocio_novo_vazio.png")], "Estabelecimento recém-criado: os pedidos de outros negócios não aparecem"),
        Figure([s3("s03_13_usuarios.png")], "Tela de usuários do administrador, com o resumo das permissões de cada perfil"),
        Figure([s3("s03_14_novo_usuario.png")], "Cadastro de um entregador com capacidade de carga"),
        Figure([s3("s03_15_usuario_cadastrado.png")], "Entregador cadastrado e listado com disponibilidade e capacidade"),
        H("18  Controle de perfis"),
        P("Os três perfis previstos no RF01 foram implementados. A verificação acontece na API, em cada rota, e não apenas na interface: mesmo que alguém chame a API diretamente, a ação é negada para perfis sem permissão. O perfil é lido do banco a cada requisição, então uma alteração feita pelo administrador vale imediatamente."),
        Table(
            ["Ação", "Administrador", "Atendente", "Entregador"],
            [
                ["Entrar no sistema", "Sim", "Sim", "Sim"],
                ["Cadastrar, editar e desativar usuários", "Sim", "Não", "Não"],
                ["Consultar pedidos", "Todos do estabelecimento", "Todos do estabelecimento", "Só os atribuídos a ele"],
                ["Cadastrar e editar pedidos", "Sim", "Sim", "Não"],
                ["Atribuir entregador ao pedido", "Sim", "Sim", "Não"],
                ["Excluir pedidos", "Sim", "Não", "Não"],
                ["Atualizar a situação", "Qualquer situação", "Qualquer situação", "De Atribuído para Em rota e de Em rota para Entregue"],
            ],
            [5.0, 3.6, 3.6, 5.0],
        ),
        Figure([s3("s03_16_atendente.png")], "Visão do atendente: sem a aba Usuários e sem o botão Excluir"),
        Figure(
            [s3("s03_17_entregador_celular.png"), s3("s03_18_entregador_status.png")],
            "Visão do entregador no celular: apenas os próprios pedidos e o botão para avançar a entrega",
            36,
        ),
        Code(evidence("api_controle_acesso.txt"), "Respostas da API para cada perfil, obtidas com chamadas diretas"),
        H("19  CRUD principal: pedidos"),
        P("O pedido é a entidade principal do RotaCerta porque concentra o que será roteirizado: cliente, destino, carga, prioridade, janela de horário, entregador e situação. O CRUD usa as tabelas orders e customers do modelo da Sprint 02. O cliente é informado no próprio formulário do pedido; a API reaproveita o cadastro quando o mesmo nome e telefone já existem e cria um registro em customers quando não existem. O local de entrega é marcado com um clique no mapa, que preenche latitude e longitude."),
        Table(
            ["Operação", "Rota da API", "Regra de negócio"],
            [
                ["Cadastrar", "POST /orders", "Coordenadas obrigatórias, fim da janela posterior ao início; com entregador o pedido já nasce Atribuído"],
                ["Consultar", "GET /orders e GET /orders/{id}", "Lista com filtro por situação; o entregador vê só os próprios pedidos"],
                ["Atualizar", "PUT /orders/{id} e PATCH /orders/{id}/status", "Pedidos entregues ou cancelados não podem ser editados; Em rota exige entregador"],
                ["Excluir", "DELETE /orders/{id}", "Somente administrador; pedido que já faz parte de uma rota é protegido pela chave estrangeira"],
            ],
            [2.6, 5.4, 9.2],
        ),
        Figure([s3("s03_07_pedidos_admin.png")], "Lista de pedidos do administrador com filtros por situação"),
        Figure([s3("s03_08_novo_pedido.png")], "Cadastro de pedido com cliente, local no mapa, prioridade, janela e entregador"),
        Figure([s3("s03_09_pedido_cadastrado.png")], "Pedido cadastrado aparecendo no topo da lista"),
        Figure([s3("s03_10_editar_pedido.png")], "Edição do pedido: prioridade, horário e entregador alterados"),
        Code(evidence("sql_pedido_editado.txt"), "Pedido gravado no banco com os valores da edição"),
        Figure([s3("s03_12_pedido_excluido.png")], "Pedido excluído pelo administrador"),
        Code(evidence("sql_pedido_excluido.txt"), "O pedido excluído não existe mais no banco"),
        H("20  Execução local"),
        H("Com Docker", 3),
        Numbered([
            "Instalar o Docker Desktop.",
            "Na raiz do repositório, copiar .env.example para .env.",
            "Executar docker compose up --build.",
            "Se já existir um PostgreSQL usando a porta 5432, alterar POSTGRES_PORT no .env para 5433.",
        ]),
        H("Sem Docker", 3),
        Numbered([
            "Instalar Python 3.12 ou superior, Node.js 20 ou superior e PostgreSQL 16 ou superior.",
            "No PostgreSQL, criar o usuário e o banco: CREATE ROLE rotacerta LOGIN PASSWORD 'rotacerta_dev'; CREATE DATABASE rotacerta OWNER rotacerta;",
            "Na pasta backend, criar o ambiente virtual, instalar requirements.txt e executar uvicorn app.main:app --reload. Na primeira execução a API cria as tabelas e os dados de demonstração.",
            "Na pasta frontend, executar npm install e npm run dev.",
            "Para outra porta, usuário ou senha do banco, copiar backend/.env.example para backend/.env e ajustar DATABASE_URL.",
        ]),
        Table(
            ["Serviço", "Endereço"],
            [
                ["Sistema", "http://localhost:5173"],
                ["Documentação interativa da API", "http://localhost:8000/docs"],
                ["Verificação da API e do banco", "http://localhost:8000/health"],
            ],
            [7.0, 10.2],
            9,
        ),
        Table(
            ["Perfil", "E-mail", "Senha"],
            [
                ["Administrador", "admin@rotacerta.com.br", "rotacerta123"],
                ["Atendente", "atendente@rotacerta.com.br", "rotacerta123"],
                ["Entregador", "entregador@rotacerta.com.br", "rotacerta123"],
            ],
            [4.0, 8.2, 5.0],
            9,
        ),
        P("As contas acima são criadas automaticamente apenas quando o banco não possui nenhum estabelecimento, para facilitar a demonstração. O passo a passo completo também está no README do repositório."),
        H("21  Testes automatizados"),
        P("Os testes da API usam um banco SQLite em memória, recriado a cada teste, e cobrem login, sessão, hash de senha, cadastro de usuários, isolamento entre estabelecimentos, permissões de cada perfil, CRUD de pedidos e regras de situação. Os testes do otimizador continuam validando os modos sequencial e paralelo."),
        Code(evidence("testes.txt"), "Execução da suíte de testes"),
        H("22  Ajustes no planejamento e na modelagem"),
        H("Modelagem", 3),
        P("Nenhuma tabela, coluna ou restrição do modelo relacional da Sprint 02 precisou ser alterada. A implementação tomou três decisões que não mudam o modelo:"),
        Bullets([
            "O cliente é cadastrado a partir do formulário do pedido, sem uma tela própria nesta etapa, reaproveitando a tabela customers.",
            "Usuários são desativados em vez de excluídos, porque pedidos e rotas referenciam entregadores por chave estrangeira.",
            "As coordenadas do pedido são obtidas com um clique no mapa Leaflet; a busca automática pelo endereço fica para a próxima etapa.",
        ]),
        H("Planejamento", 3),
        P("A Sprint 03 da disciplina exigiu login, perfis, cadastro de usuários e o CRUD principal, que no cronograma da Sprint 02 estavam previstos para as Sprints 04 e 05. Esses itens foram antecipados. As etapas seguintes foram reorganizadas para manter a otimização sequencial e paralela e os experimentos antes do período de relatórios e da banca."),
        Table(
            ["Sprint", "Período", "Entrega principal", "Situação"],
            [
                ["1", "03/09 a 05/09", "Planejamento, requisitos e backlog", "Concluída"],
                ["2", "06/09 a 19/09", "Arquitetura, modelos, protótipos, banco e estrutura", "Concluída"],
                ["3", "20/09 a 26/09", "Banco conectado, login, perfis, usuários e CRUD de pedidos", "Concluída"],
                ["4", "27/09 a 03/10", "Entregadores, disponibilidade, clientes e busca de endereço", "Planejada"],
                ["5", "04/10 a 10/10", "Atribuição de pedidos e visualização das rotas no mapa", "Planejada"],
                ["6", "11/10 a 17/10", "Vizinho mais próximo integrado aos pedidos do dia", "Planejada"],
                ["7", "18/10 a 24/10", "2-opt, persistência das rotas e linha de base", "Planejada"],
                ["8", "25/10 a 31/10", "Paralelização para múltiplos entregadores", "Planejada"],
                ["9", "01/11 a 07/11", "Experimentos sequencial versus paralelo", "Planejada"],
                ["10", "08/11 a 14/11", "Histórico, relatórios e painel do dia", "Planejada"],
                ["11", "15/11 a 21/11", "Testes, usabilidade e segurança", "Planejada"],
                ["12", "22/11 a 28/11", "Documentação e ajustes", "Planejada"],
                ["Final", "29/11 a 05/12", "Correções, vídeos e preparação para a banca", "Planejada"],
            ],
            [1.6, 3.0, 9.6, 3.0],
        ),
        H("23  Repositório GitHub"),
        P(f"Repositório oficial. {REPOSITORY}", "Repositório oficial."),
        Table(
            ["Branch", "Finalidade"],
            [
                ["master", "Versão estável e entregas aprovadas"],
                ["sprint/02-arquitetura-modelagem", "Integração da Sprint 02"],
                ["sprint/03-estrutura-funcional", "Integração da Sprint 03"],
                ["dev/alvaro, dev/arthur, dev/vinicius", "Ambientes pessoais dos integrantes"],
            ],
            [7.0, 10.2],
            9,
        ),
        Table(
            ["Caminho", "Novidade da Sprint 03"],
            [
                ["backend/app/routers/", "Rotas de autenticação, usuários e pedidos"],
                ["backend/app/models.py e database.py", "Conexão e mapeamento do PostgreSQL"],
                ["backend/tests/test_api.py", "Testes da API"],
                ["frontend/src/pages/", "Telas de login, pedidos e usuários"],
                ["docs/sprint-03/", "Capturas de tela e saídas usadas como evidência"],
                ["tools/documentation/", "Captura das evidências e geração deste documento"],
            ],
            [6.5, 10.7],
            9,
        ),
        H("24  Dificuldades encontradas e próximos passos"),
        H("Dificuldades", 3),
        Table(
            ["Dificuldade", "Como foi tratada"],
            [
                ["Docker indisponível em máquina de desenvolvimento", "Execução sem Docker documentada; a API cria as tabelas a partir do schema.sql quando o banco está vazio"],
                ["PostgreSQL já instalado ocupando a porta 5432", "Porta do banco no Compose passou a ser configurável por POSTGRES_PORT"],
                ["psycopg 3.2.9 sem pacote para o Python 3.14", "Atualização para psycopg 3.3.6, compatível com Python 3.12 a 3.14"],
                ["Pedidos exigem latitude e longitude", "Mapa no formulário para marcar o ponto; busca automática pelo endereço planejada"],
                ["Paralelização mais lenta que a versão sequencial em entradas pequenas", "Criar processos custa mais que o cálculo com poucos pedidos; os experimentos usarão cenários de 50 a 500 pedidos"],
            ],
            [6.2, 11.0],
        ),
        H("Próximos passos", 3),
        Bullets([
            "Gerenciar entregadores e disponibilidade em tela própria e buscar coordenadas a partir do endereço.",
            "Integrar o otimizador aos pedidos do dia e gravar rotas em routes e route_stops.",
            "Registrar as execuções sequenciais e paralelas em optimization_runs.",
            "Distribuir as tarefas entre os integrantes com pull requests revisados, registrando a contribuição de cada um no GitHub.",
        ]),
        H("Situação da Sprint 03", 2),
        Table(
            ["Entrega obrigatória", "Situação", "Evidência"],
            [
                ["Banco de dados conectado", "Concluído", "Seção 15: /health, consulta às tabelas e documentação da API"],
                ["Login funcional", "Concluído", "Seção 16: tela de login, erro de senha e hash no banco"],
                ["Cadastro de usuários", "Concluído", "Seção 17: cadastro de estabelecimento e de usuários"],
                ["Controle de perfis", "Concluído", "Seção 18: matriz de permissões, telas por perfil e respostas 401 e 403"],
                ["CRUD principal", "Concluído", "Seção 19: cadastro, consulta, edição e exclusão de pedidos no banco"],
                ["Deploy local", "Concluído", "Seção 20: Docker Compose ou execução manual"],
                ["Repositório atualizado", "Concluído", "Seção 23: branch da Sprint 03"],
            ],
            [4.2, 2.6, 10.4],
        ),
    )


def sprint04(doc: Document) -> None:
    doc.add(
        PageBreak(),
        H("Parte IV  Sprint 04 Primeiro módulo completo", 1),
        P("A Sprint 04 consolida o primeiro módulo completo do RotaCerta: Pedidos e entregas. O módulo cobre o ciclo inteiro de um pedido, do cadastro do cliente até a confirmação da entrega pelo entregador, com dados gravados no PostgreSQL, validações na interface e na API, mensagens de erro que orientam o usuário e navegação por páginas com endereço próprio."),
        H("25  Módulo implementado: Pedidos e entregas"),
        P("O módulo atende aos casos de uso UC01 (cadastrar pedido), UC04 (visualizar a rota atribuída) e UC05 (atualizar o status da entrega) e aos requisitos RF02, RF05, RF07 e RF08. Ele usa a estrutura da Sprint 03 (login, perfis e usuários) e acrescenta o cadastro de clientes, a busca de endereço, as regras de atribuição e o histórico de cada pedido."),
        Table(
            ["Funcionalidade", "Perfis", "Tela"],
            [
                ["Cadastro, busca, edição e exclusão de clientes", "Administrador e atendente", "/clientes"],
                ["Cadastro de pedido com cliente, endereço, local no mapa, peso, prioridade e horário", "Administrador e atendente", "/pedidos/novo"],
                ["Busca de endereço no OpenStreetMap com preferência para a região do ponto de saída", "Administrador e atendente", "/pedidos/novo"],
                ["Atribuição a um entregador respeitando disponibilidade e capacidade de carga", "Administrador e atendente", "/pedidos/novo e /pedidos/:id/editar"],
                ["Lista de pedidos com busca e filtro por situação guardados no endereço da página", "Administrador e atendente", "/pedidos"],
                ["Detalhe do pedido com mapa e histórico de cada mudança de situação", "Todos", "/pedidos/:id"],
                ["Cancelamento com motivo obrigatório, reabertura e exclusão com confirmação", "Administrador e atendente; exclusão só administrador", "/pedidos/:id"],
                ["Entregas do dia no celular: sair para entrega, confirmar entrega e informar disponibilidade", "Entregador", "/entregas"],
            ],
            [7.8, 4.4, 5.0],
        ),
        H("Fluxo completo de utilização", 3),
        Numbered([
            "O atendente abre Novo pedido e, como a cliente ainda não tem cadastro, usa Cadastrar cliente. Ao salvar, o sistema volta ao pedido com a cliente já escolhida.",
            "O atendente digita o endereço e usa Buscar endereço. O ponto escolhido aparece no mapa e pode ser ajustado com um clique.",
            "O atendente informa peso, prioridade e horário e escolhe o entregador. A lista mostra a carga atual de cada um e o sistema impede ultrapassar a capacidade.",
            "Ao salvar, o pedido é gravado como Atribuído e o sistema abre a página do pedido com o histórico.",
            "No celular, a entregadora vê o pedido em Minhas entregas, toca em Sair para entrega e, ao chegar, em Confirmar entrega.",
            "O administrador acompanha o histórico completo, com quem fez cada mudança e quando, e pode cancelar, reabrir ou excluir pedidos.",
        ]),
        H("26  Evidências do módulo funcionando"),
        P("As figuras a seguir seguem o fluxo acima com a cliente Luciana Barros, cadastrada durante a captura. Os dados de demonstração (Farmácia Boa Saúde, com dois entregadores e seis pedidos) são criados automaticamente quando o banco está vazio."),
        Figure([s4("s04_02_pedidos.png")], "Lista de pedidos do atendente, com busca e filtros por situação"),
        Figure([s4("s04_05_busca_endereco.png")], "Pedido com a cliente recém-cadastrada já escolhida e o endereço encontrado pela busca"),
        Figure([s4("s04_07_pedido_criado.png")], "Página do pedido criado, com dados, mapa e histórico"),
        Figure(
            [s4("s04_10_entregas_celular.png"), s4("s04_12_entregador_detalhe.png")],
            "Entregadora no celular: lista de entregas e confirmação da entrega na página do pedido",
            36,
        ),
        Figure([s4("s04_13_historico_completo.png")], "Visão do administrador: o histórico registra cadastro, atribuição, saída e entrega com autor e horário"),
        Figure([s4("s04_16_pedido_cancelado.png")], "Pedido cancelado com o motivo registrado no histórico"),
        H("27  Persistência de dados"),
        P("Todos os dados do módulo ficam no PostgreSQL: clientes na tabela customers, pedidos em orders e cada mudança de situação na nova tabela order_status_history. Para demonstrar a persistência, depois do fluxo completo a API e o próprio PostgreSQL foram desligados. Com a API fora do ar, o sistema informou que o servidor não respondia. Em seguida o banco e a API foram religados e o mesmo pedido foi consultado, com os mesmos dados e o mesmo histórico."),
        Code(evidence04("persistencia.txt"), "Registro com horário de cada etapa do teste de persistência"),
        Code(evidence04("sql_pedido_criado.txt"), "Pedido e histórico no banco logo depois do cadastro pela interface"),
        Figure([s4("s04_20_api_fora.png")], "Tentativa de entrar com a API desligada: o sistema explica o problema em vez de travar"),
        Code(evidence04("sql_apos_reinicio.txt"), "O mesmo pedido consultado no banco depois de desligar e religar o PostgreSQL e a API"),
        Figure([s4("s04_21_apos_reinicio.png")], "Página do pedido depois da reabertura do sistema, com o histórico preservado"),
        P("Na execução com Docker Compose os dados ficam no volume postgres_data, que é mantido ao parar e subir os contêineres. Apenas docker compose down -v apaga o volume e recomeça o banco do zero."),
        H("28  Validações"),
        P("As validações existem em duas camadas. A interface verifica os campos antes de enviar, mostra a mensagem junto ao campo e leva o foco ao primeiro campo com problema. A API repete as mesmas regras e aplica as regras de negócio que dependem do banco, de modo que nenhum dado inválido é gravado mesmo que alguém chame a API diretamente."),
        Table(
            ["Campo ou regra", "Onde", "Regra"],
            [
                ["E-mail e senha no login", "Interface", "Obrigatórios; e-mail em formato válido"],
                ["Senha de novos usuários", "Interface e API", "Mínimo de 8 caracteres, com letras e números"],
                ["Nome de cliente e de usuário", "Interface e API", "Obrigatório, com pelo menos duas letras; espaços extras são removidos"],
                ["Telefone do cliente", "Interface e API", "Opcional; se informado, precisa ter DDD. É gravado no formato (81) 98800-1001"],
                ["Cliente duplicado", "API", "Não permite outro cliente com o mesmo nome e telefone no estabelecimento"],
                ["Cliente, endereço e local do pedido", "Interface e API", "Cliente do próprio estabelecimento; endereço com pelo menos 5 caracteres; coordenadas marcadas no mapa"],
                ["Peso do pedido", "Interface e API", "Maior que zero e até 500 kg"],
                ["Janela de entrega", "Interface e API", "Fim depois do início; no cadastro, o horário final não pode estar no passado"],
                ["Raio de entrega", "API", "O local precisa estar a até 30 km do ponto de saída do estabelecimento"],
                ["Capacidade do entregador", "Interface e API", "A soma dos pedidos atribuídos e em rota não pode passar da capacidade de carga"],
                ["Disponibilidade", "Interface e API", "Entregador fora de serviço não recebe pedidos e só fica fora de serviço sem entregas em rota"],
                ["Mudança de situação", "API", "O entregador só avança de Atribuído para Em rota e para Entregue; pedidos entregues ou cancelados não são editados"],
                ["Motivo do cancelamento", "Interface", "Obrigatório, com pelo menos 5 caracteres, e registrado no histórico"],
            ],
            [4.6, 3.0, 9.6],
        ),
        Figure([s4("s04_01_login_validacao.png")], "Login enviado sem preencher os campos"),
        Figure([s4("s04_03_pedido_validacao.png")], "Pedido enviado vazio: cada campo obrigatório indica o que falta e o mapa fica destacado"),
        Figure([s4("s04_04_cliente_validacao.png")], "Telefone sem DDD recusado no cadastro de cliente"),
        Figure([s4("s04_06_capacidade.png")], "Pedido de 30 kg recusado para uma entregadora com capacidade de 25 kg"),
        Figure([s4("s04_15_cancelamento_validacao.png")], "Cancelamento sem motivo não é aceito"),
        Code(evidence04("api_validacoes.txt"), "Respostas da API às mesmas regras quando chamada diretamente"),
        P("As mensagens de validação padrão da API aparecem em inglês quando ela é chamada diretamente, como no quadro acima. A interface traduz essas mensagens e as exibe em português junto ao campo correspondente."),
        H("29  Mensagens de erro"),
        P("Cada tipo de falha tem uma mensagem própria, escrita para orientar o próximo passo do usuário. Nenhuma tela mostra erros técnicos, códigos internos ou páginas em branco."),
        Table(
            ["Situação", "Como o sistema responde"],
            [
                ["Campo inválido", "Mensagem abaixo do campo, borda vermelha, resumo no topo do formulário e foco no primeiro campo com erro"],
                ["Regra de negócio recusada pela API", "A mensagem da API aparece junto ao campo relacionado, como entregador, local no mapa ou horário"],
                ["Ação que não pode ser concluída", "Aviso vermelho temporário explicando o motivo, como cliente com pedidos ou entrega em rota"],
                ["Ação irreversível", "Janela de confirmação antes de excluir pedidos, clientes ou desativar usuários"],
                ["Página de outro perfil", "Tela Acesso negado indicando quais perfis usam a página"],
                ["Endereço inexistente", "Tela Página não encontrada com o endereço digitado e link para o início"],
                ["API ou internet fora do ar", "Mensagem de servidor indisponível e, nas listas, a opção Tentar de novo"],
                ["Busca de endereço indisponível", "Orientação para marcar o local direto no mapa"],
                ["Sessão expirada", "Volta para o login com o aviso Sua sessão expirou"],
                ["Falha inesperada ao exibir uma tela", "Tela de erro com o botão Recarregar a página, sem afetar o restante do sistema"],
            ],
            [5.0, 12.2],
        ),
        Figure([s4("s04_11_indisponivel_erro.png")], "Entregadora com entrega em rota tentando ficar fora de serviço", 40),
        Figure([s4("s04_17_confirmacao.png")], "Confirmação antes de excluir um cliente"),
        Figure([s4("s04_18_cliente_com_pedidos.png")], "Cliente com pedidos não pode ser excluído"),
        Figure([s4("s04_19_raio_entrega.png")], "Local de entrega fora do raio de 30 km, recusado pela API e indicado no mapa"),
        Figure([s4("s04_08_acesso_negado.png")], "Atendente tentando abrir a página de usuários"),
        Figure([s4("s04_09_pagina_inexistente.png")], "Endereço que não existe no sistema"),
        H("30  Navegação entre as telas"),
        P("O frontend passou a usar o React Router. Cada tela tem um endereço próprio, o botão Voltar do navegador funciona, a busca e o filtro da lista ficam no endereço e podem ser compartilhados, e as páginas internas mostram o caminho de navegação (Pedidos / Pedido #7). As rotas são protegidas: sem login o sistema leva para a tela de entrada e, depois do login, abre a página pedida originalmente."),
        Figure([s4("s04_00_mapa_navegacao.png")], "Mapa de navegação entre as telas do módulo"),
        Table(
            ["Endereço", "Tela", "Perfis"],
            [
                ["/login e /cadastro", "Entrada e cadastro de novo negócio", "Público"],
                ["/pedidos", "Lista de pedidos com busca e filtros", "Administrador e atendente"],
                ["/pedidos/novo", "Cadastro de pedido", "Administrador e atendente"],
                ["/pedidos/:id", "Detalhe e histórico do pedido", "Todos; o entregador só abre os seus"],
                ["/pedidos/:id/editar", "Edição do pedido", "Administrador e atendente"],
                ["/clientes, /clientes/novo e /clientes/:id/editar", "Clientes", "Administrador e atendente"],
                ["/entregas", "Entregas do entregador", "Entregador"],
                ["/usuarios, /usuarios/novo e /usuarios/:id/editar", "Usuários", "Administrador"],
            ],
            [6.4, 5.8, 5.0],
        ),
        Figure([s4("s04_14_filtro_url.png")], "Filtro por situação guardado no endereço /pedidos?situacao=PENDING"),
        Code(evidence04("navegacao.txt"), "Retorno à página pedida depois do login"),
        H("31  Commits organizados e repositório"),
        P(f"Repositório oficial. {REPOSITORY}", "Repositório oficial."),
        P("A Sprint 04 foi desenvolvida na branch sprint/04-modulo-pedidos-entregas, criada a partir da master com a Sprint 03 já integrada pelo Pull Request número 2. Cada etapa foi registrada em um commit próprio no momento em que ficou pronta e testada, seguindo o padrão do CONTRIBUTING.md (feat, fix, docs)."),
        Code(evidence04("commits.txt"), "Commits da Sprint 04 em ordem cronológica"),
        H("32  Ajustes no planejamento, na arquitetura e na modelagem"),
        H("Modelagem", 3),
        P("Foi criada a tabela order_status_history (id, order_id, status, note, changed_by, changed_at), ligada a orders com exclusão em cascata e a users pelo autor da mudança. Ela atende ao RF08, manter o histórico das entregas, que o modelo da Sprint 02 cobria apenas pela situação atual do pedido. Bancos criados antes desta versão recebem a tabela automaticamente quando a API inicia. As demais tabelas não mudaram."),
        H("Arquitetura e API", 3),
        Bullets([
            "O pedido passou a referenciar um cliente cadastrado (customer_id), em vez de criar o cliente pelo nome digitado, o que evita cadastros duplicados.",
            "Novas rotas: /customers, /geocode, /couriers/me/availability e o histórico em GET /orders/{id}.",
            "As regras de atribuição ficaram concentradas no módulo backend/app/rules.py, usado pelas rotas de pedidos e de entregadores.",
            "A busca de endereço usa o serviço Nominatim do OpenStreetMap por meio da API, que identifica o sistema como exige a política do serviço. Se ele estiver indisponível, o usuário marca o ponto no mapa.",
            "No frontend, o React Router substituiu a troca de telas por abas, e os formulários laterais viraram páginas próprias.",
        ]),
        H("Planejamento", 3),
        P("O módulo escolhido reúne o que o cronograma previa para as Sprints 04 e 05 (clientes, busca de endereço, entregadores e disponibilidade) e antecipa o histórico de entregas. As etapas seguintes continuam priorizando o componente de otimização."),
        Table(
            ["Sprint", "Período", "Entrega principal", "Situação"],
            [
                ["1", "03/09 a 05/09", "Planejamento, requisitos e backlog", "Concluída"],
                ["2", "06/09 a 19/09", "Arquitetura, modelos, protótipos, banco e estrutura", "Concluída"],
                ["3", "20/09 a 26/09", "Banco conectado, login, perfis, usuários e CRUD de pedidos", "Concluída"],
                ["4", "27/09 a 03/10", "Módulo Pedidos e entregas completo", "Concluída"],
                ["5", "04/10 a 10/10", "Rotas do dia no mapa e atribuição automática de pedidos", "Planejada"],
                ["6", "11/10 a 17/10", "Vizinho mais próximo integrado aos pedidos do dia", "Planejada"],
                ["7", "18/10 a 24/10", "2-opt, gravação das rotas e linha de base", "Planejada"],
                ["8", "25/10 a 31/10", "Paralelização para múltiplos entregadores", "Planejada"],
                ["9", "01/11 a 07/11", "Experimentos sequencial versus paralelo", "Planejada"],
                ["10", "08/11 a 14/11", "Relatórios e painel do dia", "Planejada"],
                ["11", "15/11 a 21/11", "Testes, usabilidade e segurança", "Planejada"],
                ["12", "22/11 a 28/11", "Documentação e ajustes", "Planejada"],
                ["Final", "29/11 a 05/12", "Correções, vídeos e preparação para a banca", "Planejada"],
            ],
            [1.6, 3.0, 9.6, 3.0],
        ),
        H("33  Testes automatizados"),
        P("A suíte passou de 21 para 48 testes. Os novos testes cobrem clientes, validações de telefone, nome e senha, regras de capacidade, disponibilidade, raio e horário, o histórico de situação e a busca de endereço, esta sem acessar a internet."),
        Code(evidence04("testes.txt"), "Execução da suíte de testes"),
        H("34  Dificuldades encontradas e próximos passos"),
        H("Dificuldades", 3),
        Table(
            ["Dificuldade", "Como foi tratada"],
            [
                ["O mapa escapava do formulário quando o campo entrava em estado de erro", "O React substituía as classes que o Leaflet adiciona ao elemento do mapa. O mapa passou a ficar em um elemento interno que o React não altera"],
                ["Dependência de um serviço externo para buscar endereços", "A API identifica o sistema, limita os resultados à região e, se o serviço falhar, orienta a marcar o ponto no mapa. Os testes simulam o serviço"],
                ["Bancos da Sprint 03 sem a tabela de histórico", "A API cria a tabela ao iniciar quando ela não existe, sem apagar dados"],
                ["Manter commits organizados durante o desenvolvimento", "O trabalho foi dividido em etapas testadas separadamente e cada etapa gerou um commit próprio"],
            ],
            [6.2, 11.0],
        ),
        H("Próximos passos", 3),
        Bullets([
            "Mostrar no mapa as rotas do dia de cada entregador e sugerir a atribuição automática dos pedidos pendentes.",
            "Integrar o otimizador sequencial e paralelo aos pedidos reais e gravar as rotas em routes e route_stops.",
            "Registrar as execuções em optimization_runs para os experimentos de desempenho.",
            "Distribuir as próximas tarefas entre os integrantes, cada um com os próprios commits e pull requests revisados por outro colega.",
        ]),
        H("Situação da Sprint 04", 2),
        Table(
            ["Entrega obrigatória", "Situação", "Evidência"],
            [
                ["Primeiro módulo totalmente funcional", "Concluído", "Seções 25 e 26: fluxo completo de Pedidos e entregas"],
                ["Persistência de dados", "Concluído", "Seção 27: dados preservados depois de desligar e religar banco e API"],
                ["Validações", "Concluído", "Seção 28: regras na interface e na API"],
                ["Mensagens de erro", "Concluído", "Seção 29: validação, regras, permissões, servidor indisponível"],
                ["Navegação entre telas", "Concluído", "Seção 30: rotas, mapa de navegação e caminho da página"],
                ["Commits organizados", "Concluído", "Seção 31: um commit por etapa na branch da Sprint 04"],
            ],
            [4.4, 2.6, 10.2],
        ),
    )


TEST_COUNT = 76

EXPERIMENT_ANALYSIS = (
    "Três comportamentos aparecem nos experimentos. Primeiro, com os pedidos reais da demonstração (três paradas), "
    "todos os modos paralelos ficam mais lentos que o sequencial: o cálculo leva cerca de 0,1 ms, menos que o custo fixo "
    "de enviar os dados a outro processo (cerca de 0,5 ms) ou de copiá-los para a GPU e lançar o kernel (cerca de 1 ms). "
    "Para a operação diária de um pequeno negócio o modo sequencial já basta; o paralelismo compensa quando o problema "
    "cresce, como em várias filiais, muitos entregadores ou simulações. Segundo, na CPU a melhor eficiência ficou entre "
    "2 e 4 processos (90% e 74% na instância de 16 rotas com 100 paradas) e o speedup parou de crescer acima de 8 "
    "processos, porque o processador tem 8 núcleos físicos e os 16 núcleos lógicos dividem os mesmos recursos. Nas duas "
    "instâncias maiores o ganho da CPU foi menor e até o pool com 1 processo ficou mais lento que o sequencial; num script "
    "isolado, fora da API, os dois tempos são iguais (cerca de 200 ms), por isso a hipótese é a disputa de CPU com outros "
    "programas abertos durante a medição, registrada como pendência na seção 42. Terceiro, o ganho da GPU cresce com o "
    "tamanho do problema, de 1,8 vezes com 8 rotas de 50 paradas para 24,5 vezes com 32 rotas de 300 paradas: mais rotas "
    "ocupam mais dos 28 multiprocessadores da RTX 3060 e mais paradas significam mais trocas do 2-opt avaliadas ao mesmo "
    "tempo. Em todas as medições os três modos chegaram exatamente às mesmas rotas e à mesma distância total."
)


def sprint05(doc: Document) -> None:
    doc.add(
        PageBreak(),
        H("Parte V  Sprint 05 Segundo módulo funcionando", 1),
        P("A Sprint 05 entrega o segundo módulo do RotaCerta: Roteirização e desempenho. Seguindo a orientação de avançar de forma objetiva para o núcleo de otimização, o módulo liga o otimizador aos pedidos reais gravados no banco, mostra no mapa a rota de cada entregador, permite que o entregador inicie e conclua a rota pelo celular e compara a execução sequencial, a paralela em CPU e a paralela em GPU com CUDA, registrando tempo, speedup e eficiência."),
        H("35  Módulo implementado: Roteirização e desempenho"),
        P("O módulo atende aos casos de uso UC03 (gerar rotas otimizadas), UC04 (visualizar a rota do dia) e UC05 (atualizar o andamento da entrega) e aos requisitos RF04 (rota otimizada por entregador), RF06 (rota como lista ordenada e mapa), RF07 e RF08 (situação e histórico), além do RNF01 (desempenho com processamento paralelo). Ele usa a estrutura e o primeiro módulo das sprints anteriores: login, perfis, clientes, pedidos e atribuição."),
        Table(
            ["Funcionalidade", "Perfis", "Tela ou rota da API"],
            [
                ["Geração das rotas do dia a partir dos pedidos atribuídos, no modo sequencial, paralelo em CPU ou GPU (CUDA)", "Administrador e atendente", "/rotas  |  POST /routes/generate"],
                ["Mapa com o ponto de saída, o trajeto de cada entregador e as paradas numeradas; lista ordenada com horário estimado", "Administrador e atendente", "/rotas  |  GET /routes"],
                ["Cadastro do ponto de saída no mapa, com busca de endereço", "Administrador", "/rotas  |  PUT /establishment/depot"],
                ["Rota do entregador no celular, com mapa, próxima parada e botão Iniciar rota", "Entregador", "/entregas  |  GET /routes/me e PATCH /routes/{id}/start"],
                ["Conclusão automática da parada e da rota conforme as entregas são confirmadas", "Entregador", "/entregas  |  PATCH /orders/{id}/status"],
                ["Comparação de desempenho com tempo, speedup, eficiência e gráfico, em instância sintética ou nos pedidos reais", "Administrador", "/desempenho  |  POST /optimizer/benchmark"],
                ["Histórico das comparações e das gerações de rotas gravado no banco", "Administrador", "/desempenho  |  GET /optimizer/runs"],
            ],
            [8.4, 3.6, 5.2],
        ),
        H("Fluxo completo de utilização", 3),
        Numbered([
            "O atendente cadastra e atribui os pedidos aos entregadores, como no primeiro módulo.",
            "Na tela Rotas, ele escolhe o modo de execução e clica em Gerar rotas. A API agrupa os pedidos atribuídos por entregador, executa o otimizador e grava as rotas, as paradas e a execução no banco.",
            "O mapa mostra o trajeto de cada entregador, saindo e voltando ao ponto de saída, e a lista traz a ordem das paradas, a distância, a duração e o horário estimado de cada chegada.",
            "No celular, a entregadora abre Minhas entregas, vê a rota no mapa e toca em Iniciar rota. Todos os pedidos da rota passam para Em rota, com registro no histórico.",
            "A cada entrega confirmada a parada é concluída e a próxima fica destacada. Quando todas terminam, a rota passa para Concluída.",
            "O administrador abre Desempenho, executa a comparação entre os modos e acompanha o histórico das medições.",
        ]),
        H("36  Evidências das funcionalidades"),
        P("As capturas foram feitas com o sistema em execução sobre um banco recém-criado com os dados de demonstração: a Farmácia Boa Saúde, dois entregadores e catorze pedidos, dez deles atribuídos e aguardando rota."),
        Figure([s5("s05_01_rotas_antes.png")], "Tela Rotas antes da geração: dez pedidos atribuídos aguardando rota"),
        Figure([s5("s05_02_rotas_sequencial.png")], "Rotas geradas no modo sequencial, com o trajeto de cada entregador no mapa e a lista de paradas"),
        Figure([s5("s05_03_rotas_gpu.png")], "As mesmas rotas geradas no modo GPU (CUDA): mesma ordem e mesma distância"),
        Figure([s5("s05_04_rota_destacada.png")], "Rota da Carla destacada ao clicar no cartão; as demais ficam esmaecidas"),
        Figure([s5("s05_05_ponto_saida.png")], "Cadastro do ponto de saída, de onde as rotas partem e para onde voltam"),
        Figure(
            [s5("s05_06_rota_entregador.png"), s5("s05_07_rota_em_andamento.png")],
            "Entregadora no celular: rota planejada e, depois de iniciada, com duas entregas feitas e a próxima parada destacada",
            36,
        ),
        Figure([s5("s05_08_rota_concluida.png")], "Visão do administrador com a rota da Carla concluída e a do Diego ainda planejada"),
        H("37  Integração com o banco de dados"),
        P("O módulo usa três tabelas que já constavam do modelo da Sprint 02 e ainda não tinham uso: routes guarda a rota de cada entregador (data, situação, modo de execução, distância e duração), route_stops guarda cada parada (ordem, pedido, distância desde a parada anterior, horário estimado e situação) e optimization_runs registra cada execução do otimizador com o tempo medido. As alterações no modelo estão na seção 43."),
        Code(evidence05("sql_rotas_geradas.txt"), "Rotas, paradas e execuções gravadas logo depois da geração pela interface"),
        Code(evidence05("sql_rota_concluida.txt"), "Situação da rota, das paradas e o histórico de um pedido depois das entregas"),
        P("Para demonstrar que os dados são recuperados pela aplicação, a API foi reiniciada depois do fluxo completo e a tela Rotas foi aberta de novo: as rotas, as paradas e o histórico de medições continuaram iguais, lidos do PostgreSQL."),
        Code(evidence05("persistencia.txt"), "Registro do reinício da API e contagem dos registros no banco"),
        Figure([s5("s05_12_rotas_apos_reinicio.png")], "Tela Rotas depois de reiniciar a API, com as mesmas rotas"),
        H("38  Regras de negócio"),
        P("As regras abaixo foram implementadas no backend (módulo backend/app/planning.py) e valem para qualquer cliente da API, não só para a interface."),
        Table(
            ["Regra", "Descrição"],
            [
                ["RN01", "Só entram na rota pedidos na situação Atribuído. Pedidos em rota, entregues ou cancelados ficam de fora."],
                ["RN02", "Entregador fora de serviço não recebe rota; a resposta informa o motivo."],
                ["RN03", "Entregador com rota em andamento não recebe outra; os pedidos novos aguardam a próxima geração."],
                ["RN04", "Gerar de novo substitui as rotas ainda planejadas. Uma rota em andamento nunca é recalculada."],
                ["RN05", "Toda rota sai do ponto de saída e volta a ele. Sem o ponto marcado no mapa, a geração é recusada e o administrador é orientado a marcá-lo."],
                ["RN06", "A ordem das paradas minimiza a distância total (vizinho mais próximo seguido de 2-opt)."],
                ["RN07", "Duração e horário de chegada são estimados com velocidade média de 25 km/h e 5 minutos de atendimento por parada, valores configuráveis."],
                ["RN08", "Iniciar a rota passa todos os pedidos atribuídos dela para Em rota e registra no histórico. Só o próprio entregador, o administrador ou o atendente iniciam; entregador fora de serviço não inicia."],
                ["RN09", "Entrega confirmada conclui a parada; pedido cancelado marca a parada como não realizada. Quando todas terminam, a rota fica Concluída."],
                ["RN10", "Trocar o entregador, mudar o endereço ou cancelar um pedido de rota planejada descarta essa rota, que precisa ser gerada de novo."],
                ["RN11", "Pedido de rota iniciada ou concluída não pode ser excluído, só cancelado. Pedido de rota planejada pode ser excluído e a rota é descartada."],
                ["RN12", "O modo GPU só fica disponível com GPU NVIDIA e CuPy instalados; sem eles, a interface desabilita a opção e a API explica o motivo."],
                ["RN13", "A comparação de desempenho é exclusiva do administrador, aceita até 12.000 paradas e usa a mediana de 1 a 5 repetições."],
                ["RN14", "Só o administrador altera o ponto de saída."],
            ],
            [1.6, 15.6],
        ),
        P("O quadro a seguir mostra as respostas da API a cada regra quando chamada diretamente, em sequência, sobre os mesmos dados."),
        Code(evidence05("regras.txt"), "Regras de negócio verificadas pela API"),
        Figure([s5("s05_09_regra_recusada.png")], "A mesma recusa vista na interface: nenhum entregador pode receber rota e o motivo de cada um é informado"),
        H("Regras alteradas em relação ao planejamento", 3),
        Table(
            ["Alteração", "Justificativa"],
            [
                ["A roteirização, prevista para as Sprints 06 a 10, foi antecipada para a Sprint 05", "Orientação da professora para avançar ao núcleo de otimização, já que a parte operacional estava madura"],
                ["Prioridade e janela de horário não alteram a ordem das paradas nesta versão", "Os três modos precisam otimizar exatamente o mesmo objetivo para a comparação de desempenho ser justa. A prioridade continua visível em cada parada; as restrições entram na próxima sprint"],
                ["A atribuição dos pedidos continua manual", "A atribuição automática, listada nos próximos passos da Sprint 04, foi adiada para concentrar a sprint no cálculo das rotas e nas medições"],
                ["O ponto de saída passou a ser marcado no mapa pelo administrador", "Estabelecimentos cadastrados pela tela de cadastro só tinham o endereço, sem coordenadas, e não conseguiriam gerar rotas"],
                ["A rota é desenhada com linhas retas entre as paradas", "O cálculo usa a distância geodésica (haversine). O traçado pelas ruas depende de um serviço externo e ficou como evolução"],
            ],
            [7.0, 10.2],
        ),
        H("39  Núcleo de otimização e paralelização com CUDA"),
        P("O problema de cada entregador é encontrar a ordem de visita das paradas que reduz a distância total, saindo e voltando ao ponto de saída. A solução usa duas heurísticas clássicas do problema do caixeiro-viajante: o vizinho mais próximo monta uma rota inicial indo sempre para a parada mais próxima ainda não visitada, e o 2-opt melhora essa rota invertendo trechos enquanto houver ganho."),
        P("Nesta sprint o núcleo foi reescrito para ficar igual nos três modos. Ele calcula uma vez a matriz de distâncias haversine entre todos os pontos e usa o 2-opt por melhor melhoria: a cada rodada avalia todas as inversões possíveis pela diferença de distância e aplica a melhor. Os três modos usam a mesma matriz, o mesmo critério de desempate (menor índice) e o mesmo limite de melhoria, por isso chegam exatamente às mesmas rotas. A comparação mede só a forma de execução, não diferenças de algoritmo."),
        Table(
            ["Modo", "Como executa", "Paralelismo"],
            [
                ["Sequencial", "Um processo calcula as rotas uma depois da outra com NumPy. É a linha de base", "Nenhum"],
                ["Paralelo em CPU", "As rotas são distribuídas em fatias entre processos de um pool que fica aberto entre as chamadas", "Um processo por núcleo, escolhido pelo usuário (1, 2, 4, 8 ou 16)"],
                ["Paralelo em GPU (CUDA)", "Um kernel CUDA próprio resolve todas as rotas num único lançamento", "Um bloco de 256 threads por rota; milhares de threads no total"],
            ],
            [3.4, 9.4, 4.4],
        ),
        H("Como o kernel CUDA funciona", 3),
        P("O kernel foi escrito em CUDA C e é compilado em tempo de execução pelo NVRTC por meio do CuPy (RawKernel). Cada bloco de threads resolve a rota de um entregador, em três etapas dentro do mesmo lançamento:"),
        Numbered([
            "Matriz de distâncias: as 256 threads do bloco dividem entre si os pares de pontos da rota e calculam a distância haversine de cada par.",
            "Vizinho mais próximo: a cada passo, cada thread procura a parada mais próxima entre as que lhe cabem, e uma redução em memória compartilhada escolhe a melhor do bloco.",
            "2-opt: a cada rodada, as threads avaliam em paralelo todas as inversões possíveis, a redução escolhe a de maior ganho e as threads invertem o trecho juntas. O laço termina quando nenhuma inversão encurta a rota.",
        ]),
        P("Para que a GPU chegue aos mesmos resultados da CPU, o kernel usa precisão dupla, o mesmo desempate por menor índice e a opção --fmad=false, que impede o compilador de fundir multiplicações e somas numa única operação com arredondamento diferente do NumPy."),
        H("40  Comparação sequencial, paralela e GPU"),
        P("A comparação segue a mesma metodologia em todas as medições: todos os modos recebem a mesma entrada; antes de medir, os processos da CPU são iniciados e o kernel CUDA é compilado, para que o tempo inclua só o cálculo e a troca de dados; cada modo é executado várias vezes e vale a mediana. O speedup é o tempo sequencial dividido pelo tempo do modo, e a eficiência é o speedup dividido pelo número de processos. Na GPU a eficiência não se aplica, porque o paralelismo é de milhares de threads e não de processos."),
        Code(evidence05("maquina.txt"), "Máquina usada nas medições"),
        Figure([s5("s05_10_desempenho.png")], "Tela Desempenho com uma comparação de 16 rotas de 200 paradas"),
        Figure([s5("s05_11_grafico_tooltip.png")], "Gráfico de speedup da CPU por número de processos, com a marca do speedup ideal e o detalhe de uma barra"),
        P("Para observar como o ganho depende do tamanho do problema, a comparação foi repetida com os pedidos reais e com instâncias sintéticas crescentes, sempre com a mesma semente."),
        Code(evidence05("experimentos.txt"), "Experimentos com tamanhos crescentes (mediana de 3 execuções; pedidos reais com 5)"),
        P(EXPERIMENT_ANALYSIS),
        Code(evidence05("sql_execucoes.txt"), "Todas as execuções registradas na tabela optimization_runs"),
        H("41  Testes das funcionalidades"),
        P(f"A suíte automatizada passou de 48 para {TEST_COUNT} testes. Os novos testes cobrem a geração das rotas e todas as regras da seção 38, a consistência entre pedidos e rotas, as permissões por perfil, a comparação de desempenho e o núcleo de otimização, incluindo a equivalência das rotas entre os modos sequencial, paralelo e GPU. O teste da GPU é executado só em máquinas com GPU NVIDIA; nas demais aparece como ignorado. Os testes da API usam um banco SQLite em memória e não dependem do PostgreSQL."),
        Table(
            ["Grupo", "O que verifica"],
            [
                ["Geração de rotas", "Uma rota por entregador com todos os pedidos atribuídos, ordem das paradas, distância, duração, gravação e listagem"],
                ["Regras de geração", "Entregador fora de serviço, rota em andamento, ponto de saída ausente, nenhum pedido aguardando, modo GPU indisponível"],
                ["Acompanhamento", "Entregador vê e inicia só a própria rota, pedidos passam para Em rota, rota concluída ao entregar todas as paradas"],
                ["Consistência", "Troca de entregador descarta a rota planejada; exclusão de pedido em rota iniciada é recusada"],
                ["Perfis", "Entregador não gera rotas, atendente não executa comparações nem altera o ponto de saída"],
                ["Comparação", "Modos e workers medidos, speedup e eficiência coerentes, mesmas rotas, limite de tamanho, pedidos reais"],
                ["Otimizador", "2-opt nunca piora a rota, distância informada confere com a rota, rotas vazias e de uma parada, CPU e GPU iguais, um único pool aberto"],
            ],
            [3.6, 13.6],
        ),
        Code(evidence05("testes.txt"), "Execução da suíte de testes"),
        P("Além da suíte automatizada, o roteiro de captura das evidências (tools/documentation/capture_sprint05_evidence.py) funciona como teste de ponta a ponta: ele executa pela interface e pela API o fluxo completo da seção 35, as regras da seção 38 e as comparações da seção 40, e falha se alguma tela ou resposta não for a esperada. Foi nessa execução que apareceram os dois bugs mais importantes da sprint."),
        H("42  Bugs encontrados e correções"),
        Table(
            ["Bug", "Como foi encontrado", "Correção", "Situação"],
            [
                ["A comparação com 16 processos derrubava o pool de processos (BrokenProcessPool) e a API respondia erro 500", "Roteiro de ponta a ponta", "Cada processo carregava o NumPy com 16 threads do OpenBLAS e os pools de 1, 2, 4, 8 e 16 processos ficavam abertos juntos, esgotando a memória. O OpenBLAS passou a usar uma thread por processo e só um pool fica aberto. Novos testes cobrem os dois pontos", "Corrigido"],
                ["Com menos rotas que processos, a comparação repetia a mesma configuração (duas linhas \"CPU 8 processos\" com 8 rotas)", "Roteiro de ponta a ponta (experimentos)", "A quantidade efetiva de processos é limitada ao número de rotas, sem repetir, com teste automatizado", "Corrigido"],
                ["A mensagem de entregador fora de serviço dizia \"os pedidos dele\" também para entregadoras", "Registro das regras (seção 38)", "Texto neutro e teste que verifica a mensagem", "Corrigido"],
                ["No celular, o mapa da rota abria sem mostrar as paradas", "Teste manual pela interface", "O enquadramento era calculado antes de o mapa ter o tamanho final; agora é refeito quando o tamanho é conhecido", "Corrigido"],
                ["Com as abas novas, o menu superior ultrapassava a largura da tela no celular", "Teste manual com 390 px de largura", "O menu passou a rolar na horizontal; a página voltou a ter exatamente a largura da tela", "Corrigido"],
                ["O rótulo do speedup ficava em cima da linha do ideal no gráfico", "Revisão visual", "O rótulo ganhou fundo próprio", "Corrigido"],
                ["Rota de uma parada aparecia como \"1 paradas\"", "Teste manual", "Plural ajustado", "Corrigido"],
                ["Um pedido reaberto depois de uma parada não realizada não poderia voltar para uma rota (order_id é único em route_stops)", "Revisão do código", "A nova parada substitui o registro antigo da rota encerrada", "Corrigido"],
                ["A troca de entregador não era detectada antes de gravar, porque o SQLAlchemy só atualiza o campo assigned_courier_id ao enviar ao banco", "Revisão do código", "A verificação passou a usar o relacionamento com o entregador", "Corrigido"],
                ["Um teste de regeração de rotas falhava só no SQLite, que reaproveita números de registros apagados", "Suíte automatizada", "O teste passou a verificar o conteúdo das rotas, não os números", "Corrigido"],
                ["As operações prontas do CuPy não compilavam na pasta do projeto, que tem acento (FÁBRICA)", "Instalação do CuPy", "O kernel foi escrito sem depender dos headers do CuPy, o que dispensa o caminho com acento", "Contornado"],
            ],
            [5.2, 3.0, 6.6, 2.4],
            8,
        ),
        H("Pendências e ações previstas", 3),
        Table(
            ["Pendência", "Ação prevista"],
            [
                ["As medições variam quando outros programas usam a CPU ou a GPU ao mesmo tempo", "Repetir os experimentos finais com a máquina dedicada, mais repetições e o registro da carga do sistema (Sprint 07)"],
                ["Dentro da API, nas instâncias maiores, o pool com 1 processo ficou mais lento que o sequencial, embora num script isolado os tempos sejam iguais", "Medir de novo com a máquina dedicada e, se a diferença continuar, medir separadamente o tempo de cálculo dentro de cada processo"],
                ["A primeira geração no modo GPU depois de iniciar a API inclui a compilação do kernel, cerca de um segundo", "Compilar o kernel quando a API inicia, em segundo plano"],
                ["O contêiner Docker da API não acessa a GPU", "Documentar a configuração com NVIDIA Container Toolkit ou manter o modo GPU só na execução sem Docker"],
            ],
            [7.6, 9.6],
        ),
        H("43  Ajustes no planejamento, na arquitetura e na modelagem"),
        H("Modelagem", 3),
        Bullets([
            "As tabelas routes, route_stops e optimization_runs, já previstas desde a Sprint 02, passaram a ser usadas e ganharam mapeamento no SQLAlchemy.",
            "A coluna execution_mode de routes e de optimization_runs passou a aceitar GPU, além de SEQUENTIAL e PARALLEL.",
            "optimization_runs ganhou as colunas purpose (geração de rotas ou comparação), input_source (pedidos reais ou instância sintética), run_group (agrupa as execuções de uma comparação), speedup, efficiency e same_routes.",
            "Bancos criados nas sprints anteriores recebem essas alterações automaticamente quando a API inicia, sem perda de dados. A atualização foi verificada num banco da Sprint 04.",
        ]),
        H("Arquitetura e API", 3),
        Bullets([
            "Novos módulos no backend: planning.py (regras das rotas), benchmark.py (metodologia da comparação), optimizer/gpu.py (kernel CUDA) e os routers routes.py e optimizer.py.",
            "Novas rotas da API: POST /routes/generate, GET /routes, GET /routes/me, PATCH /routes/{id}/start, GET e PUT /establishment/depot, GET /optimizer/capabilities, POST /optimizer/benchmark e GET /optimizer/runs.",
            "As rotas de pedidos chamam planning.py a cada mudança, para manter as rotas coerentes com os pedidos.",
            "Dependências: NumPy no backend e, opcionalmente, CuPy (backend/requirements-gpu.txt), que instala pelo pip o compilador e as bibliotecas do CUDA.",
            "Frontend: novas telas Rotas e Desempenho, o componente RouteMap e a rota do entregador dentro de Minhas entregas.",
            "A escolha entre C++ com OpenMP e CUDA, deixada em aberto no backlog da Sprint 01, foi resolvida a favor de CUDA: a equipe tem GPU NVIDIA disponível, o modelo de milhares de threads se encaixa na avaliação das trocas do 2-opt e o CuPy permite manter a API em Python.",
        ]),
        H("Planejamento", 3),
        P("Com a roteirização antecipada, o cronograma das próximas sprints foi reorganizado em torno da qualidade das rotas, dos experimentos e dos relatórios."),
        Table(
            ["Sprint", "Período", "Entrega principal", "Situação"],
            [
                ["1", "03/09 a 05/09", "Planejamento, requisitos e backlog", "Concluída"],
                ["2", "06/09 a 19/09", "Arquitetura, modelos, protótipos, banco e estrutura", "Concluída"],
                ["3", "20/09 a 26/09", "Banco conectado, login, perfis, usuários e CRUD de pedidos", "Concluída"],
                ["4", "27/09 a 28/09", "Módulo Pedidos e entregas completo", "Concluída"],
                ["5", "29/09 a 03/10", "Módulo Roteirização e desempenho com CPU paralela e GPU", "Concluída"],
                ["6", "04/10 a 10/10", "Atribuição automática dos pedidos e prioridade na ordem das paradas", "Planejada"],
                ["7", "11/10 a 17/10", "Experimentos controlados e relatório de desempenho (RF09)", "Planejada"],
                ["8", "18/10 a 24/10", "Painel administrativo do dia (RF10)", "Planejada"],
                ["9", "25/10 a 31/10", "Janelas de horário nas rotas e traçado pelas ruas", "Planejada"],
                ["10", "01/11 a 07/11", "Testes, usabilidade e segurança", "Planejada"],
                ["11", "08/11 a 21/11", "Documentação e ajustes", "Planejada"],
                ["Final", "22/11 a 05/12", "Correções, vídeos e preparação para a banca", "Planejada"],
            ],
            [1.6, 3.0, 9.6, 3.0],
        ),
        H("44  Repositório GitHub"),
        P(f"Repositório oficial. {REPOSITORY}", "Repositório oficial."),
        P("A Sprint 05 foi desenvolvida na branch sprint/05-modulo-otimizacao-rotas, criada a partir da master com a Sprint 04 integrada pelo Pull Request número 4. O passo a passo para executar o projeto foi separado no arquivo COMO-RODAR.md, com os comandos para PowerShell e Git Bash comentados um a um, e o README passou a descrever os módulos e a paralelização com CUDA."),
        Code(evidence05("commits.txt"), "Commits da Sprint 05 em ordem cronológica"),
        H("45  Dificuldades encontradas e próximos passos"),
        H("Dificuldades", 3),
        Table(
            ["Dificuldade", "Como foi tratada"],
            [
                ["O CuPy não compilava as próprias operações porque o caminho do projeto tem acento", "O kernel foi escrito em CUDA C sem depender dos headers do CuPy; as cópias de memória e o kernel funcionam em qualquer pasta"],
                ["Fazer a GPU chegar exatamente às mesmas rotas da CPU", "Precisão dupla, desempate por menor índice, mesma ordem nas somas e compilação sem fusão de multiplicação e soma"],
                ["Medir desempenho no Windows, onde criar processos é caro", "O pool de processos fica aberto e é aquecido antes de medir; as rotas são enviadas em fatias"],
                ["Variação das medições com outros programas abertos e com 8 núcleos físicos para 16 lógicos", "Mediana de várias execuções; a análise considera o limite físico de núcleos; os experimentos finais serão repetidos com a máquina dedicada"],
                ["Capturar as evidências num banco recém-criado sem apagar os dados locais de desenvolvimento", "As evidências foram capturadas num schema separado do PostgreSQL, com uma segunda instância da API e do frontend"],
            ],
            [6.2, 11.0],
        ),
        H("Próximos passos", 3),
        Bullets([
            "Atribuição automática dos pedidos pendentes respeitando disponibilidade, capacidade e raio, com a rota calculada logo em seguida.",
            "Considerar a prioridade e as janelas de horário na ordem das paradas, mantendo a equivalência entre os modos.",
            "Repetir os experimentos com a máquina dedicada e produzir o relatório de desempenho (RF09).",
            "Compilar o kernel CUDA quando a API inicia e documentar o uso da GPU com Docker.",
            "Distribuir as tarefas entre os integrantes, com pull requests revisados por outro colega.",
        ]),
        H("Situação da Sprint 05", 2),
        Table(
            ["Entrega obrigatória", "Situação", "Evidência"],
            [
                ["Segundo módulo completo", "Concluído", "Seções 35 e 36: fluxo completo de roteirização, do pedido atribuído à rota concluída"],
                ["Integração com o banco de dados", "Concluído", "Seção 37: routes, route_stops e optimization_runs, com dados recuperados depois de reiniciar a API"],
                ["Atualização das regras de negócio", "Concluído", "Seção 38: 14 regras e as alterações justificadas"],
                ["Testes das funcionalidades", "Concluído", f"Seção 41: {TEST_COUNT} testes automatizados e o roteiro de ponta a ponta"],
                ["Correção dos bugs encontrados", "Concluído", "Seção 42: bugs, correções e pendências com ação prevista"],
                ["Repositório GitHub atualizado", "Concluído", "Seção 44: branch, commits, README e COMO-RODAR.md"],
                ["Dificuldades e próximos passos", "Concluído", "Seção 45"],
            ],
            [4.4, 2.6, 10.2],
        ),
    )


def build_content() -> Document:
    doc = Document()
    sprint01(doc)
    sprint02(doc)
    sprint03(doc)
    sprint04(doc)
    sprint05(doc)
    return doc


# HTML e PDF

CSS = """
@page { size: A4; margin: 18mm 18mm 20mm 18mm; }
* { box-sizing: border-box; }
body { font-family: Arial, Helvetica, sans-serif; font-size: 10.5pt; line-height: 1.45; color: #000; margin: 0; }
.cover { height: 250mm; display: flex; flex-direction: column; align-items: center; text-align: center; page-break-after: always; }
.cover .inst { font-weight: bold; font-size: 14pt; margin-top: 4mm; }
.cover .course { margin-top: 8mm; font-size: 11pt; }
.cover .title { margin-top: 42mm; font-size: 30pt; font-weight: bold; }
.cover .subtitle { margin-top: 5mm; font-size: 16pt; font-weight: bold; }
.cover .group { margin-top: 26mm; font-size: 14pt; font-weight: bold; }
.cover .team { margin-top: 8mm; line-height: 1.9; }
.cover .year { margin-top: auto; }
h1 { font-size: 16pt; margin: 0 0 4mm; break-after: avoid; }
h2 { font-size: 13pt; margin: 6mm 0 3mm; break-after: avoid; }
h3 { font-size: 11pt; margin: 5mm 0 2mm; break-after: avoid; }
p { margin: 0 0 2.6mm; text-align: justify; }
ul, ol { margin: 0 0 3mm; padding-left: 7mm; }
li { margin-bottom: 1.2mm; }
.toc p { margin: 0 0 1.4mm; text-align: left; }
.toc p.part { font-weight: bold; margin-top: 3mm; }
.page-break { break-after: page; }
table { width: 100%; border-collapse: collapse; margin: 1mm 0 4mm; font-size: 8.5pt; break-inside: auto; }
thead { display: table-header-group; }
tr { break-inside: avoid; }
th { background: #173B67; color: #fff; padding: 1.6mm 2mm; text-align: center; border: 0.6pt solid #d9d9d9; }
td { padding: 1.5mm 2mm; border: 0.6pt solid #d9d9d9; vertical-align: middle; }
tbody tr:nth-child(even) td { background: #F5F7FA; }
figure { margin: 2mm 0 5mm; text-align: center; break-inside: avoid; }
figure .images { display: flex; justify-content: center; gap: 6mm; }
figure img { max-width: 100%; border: 0.6pt solid #d0d5dd; }
figcaption, .code-caption { font-size: 8.5pt; font-style: italic; color: #58606e; margin-top: 1.5mm; text-align: center; }
.code { break-inside: avoid; margin: 1mm 0 5mm; }
pre { margin: 0; padding: 3mm; background: #F5F7FA; border: 0.6pt solid #d9d9d9; font-family: Consolas, "Courier New", monospace; font-size: 7.2pt; line-height: 1.35; white-space: pre-wrap; word-break: break-word; }
"""


def html_text(text: str) -> str:
    return escape(text)


def render_html(doc: Document) -> str:
    parts = [
        "<!doctype html><html lang='pt-BR'><head><meta charset='utf-8'>",
        f"<title>RotaCerta {TITLE}</title><style>{CSS}</style></head><body>",
        "<section class='cover'>",
        "<div class='inst'>UNINASSAU</div>",
        "<div class='course'>Fábrica de Software e Tópicos Avançados em Ciência da Computação<br>Turma 8NA  |  Semestre 2026.2</div>",
        "<div class='title'>RotaCerta</div>",
        f"<div class='subtitle'>{TITLE}</div>",
        "<div class='group'>Grupo 04</div>",
        "<div class='team'>" + "<br>".join(f"{n}  |  {m}" for n, m in TEAM) + "</div>",
        "<div class='year'>2026</div></section>",
        "<h1>Apresentação</h1>",
        *(f"<p>{html_text(t)}</p>" for t in PRESENTATION),
        "<h1>Sumário</h1><div class='toc'>",
        *(f"<p class='{'part' if item.startswith('Parte') else ''}'>{html_text(item)}</p>" for item in TOC),
        "</div>",
    ]
    figure_number = 0
    for block in doc.blocks:
        if isinstance(block, PageBreak):
            parts.append("<div class='page-break'></div>")
        elif isinstance(block, H):
            parts.append(f"<h{block.level}>{html_text(block.text)}</h{block.level}>")
        elif isinstance(block, P):
            text = html_text(block.text)
            if block.lead and block.text.startswith(block.lead):
                text = f"<b>{html_text(block.lead)}</b>{html_text(block.text[len(block.lead):])}"
            parts.append(f"<p>{text}</p>")
        elif isinstance(block, Bullets):
            parts.append("<ul>" + "".join(f"<li>{html_text(i)}</li>" for i in block.items) + "</ul>")
        elif isinstance(block, Numbered):
            parts.append("<ol>" + "".join(f"<li>{html_text(i)}</li>" for i in block.items) + "</ol>")
        elif isinstance(block, Table):
            total = sum(block.widths)
            cols = "".join(f"<col style='width:{w / total * 100:.1f}%'>" for w in block.widths)
            head = "".join(f"<th>{html_text(h)}</th>" for h in block.headers)
            rows = "".join(
                "<tr>" + "".join(f"<td>{html_text(str(c))}</td>" for c in row) + "</tr>" for row in block.rows
            )
            parts.append(
                f"<table style='font-size:{block.font_size}pt'><colgroup>{cols}</colgroup>"
                f"<thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>"
            )
        elif isinstance(block, Figure):
            figure_number += 1
            images = "".join(f"<img src='{p.as_uri()}' style='width:{block.width}%'>" for p in block.paths)
            parts.append(
                f"<figure><div class='images'>{images}</div>"
                f"<figcaption>Figura {figure_number}  {html_text(block.caption)}</figcaption></figure>"
            )
        elif isinstance(block, Code):
            parts.append(
                f"<div class='code'><pre>{html_text(block.text)}</pre>"
                f"<div class='code-caption'>{html_text(block.caption)}</div></div>"
            )
    parts.append("</body></html>")
    return "".join(parts)


def render_pdf(html_path: Path, pdf_path: Path) -> None:
    from playwright.sync_api import sync_playwright

    footer = (
        "<div style='width:100%;font-family:Arial;font-size:8px;color:#58606e;text-align:center'>"
        f"{FOOTER}<span class='pageNumber'></span></div>"
    )
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="chrome")
        page = browser.new_page()
        page.goto(html_path.as_uri(), wait_until="networkidle")
        page.pdf(
            path=str(pdf_path),
            format="A4",
            print_background=True,
            display_header_footer=True,
            header_template="<span></span>",
            footer_template=footer,
            margin={"top": "18mm", "bottom": "20mm", "left": "18mm", "right": "18mm"},
        )
        browser.close()


# DOCX

def add_code_docx(document, text: str, caption: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    cell.width = Cm(17.2)
    set_cell_shading(cell, "F5F7FA")
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    lines = text.splitlines()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_font(run, name="Consolas", size=7)
        if index < len(lines) - 1:
            run.add_break()
    add_caption_docx(document, caption)


def add_caption_docx(document, caption: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(caption), size=9, italic=True, color=GRAY)


def render_docx(doc: Document, path: Path) -> None:
    document = setup_document()
    footer = document.sections[0].footer.paragraphs[0]
    footer.runs[0].text = FOOTER
    document.core_properties.title = f"RotaCerta {TITLE}"
    document.core_properties.keywords = "RotaCerta, Sprint 01, Sprint 02, Sprint 03, Sprint 04, Sprint 05"

    def centered(text, size, bold=False, before=0, after=6):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        set_font(p.add_run(text), size=size, bold=bold)

    centered("UNINASSAU", 14, True, after=26)
    centered("Fábrica de Software e Tópicos Avançados em Ciência da Computação\nTurma 8NA  |  Semestre 2026.2", 11, after=68)
    centered("RotaCerta", 23, True, after=12)
    centered(TITLE, 17, True, after=64)
    centered("Grupo 04", 14, True, after=12)
    for name, number in TEAM:
        centered(f"{name}  |  {number}", 11, after=5)
    centered("2026", 11, before=70)

    document.add_page_break()
    add_heading(document, "Apresentação", 1)
    for text in PRESENTATION:
        add_body(document, text)
    add_heading(document, "Sumário", 1)
    for item in TOC:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(item), bold=item.startswith("Parte"))

    figure_number = 0
    for block in doc.blocks:
        if isinstance(block, PageBreak):
            document.add_page_break()
        elif isinstance(block, H):
            add_heading(document, block.text, block.level)
        elif isinstance(block, P):
            add_body(document, block.text, block.lead)
        elif isinstance(block, Bullets):
            add_bullets(document, block.items)
        elif isinstance(block, Numbered):
            add_numbered(document, block.items)
        elif isinstance(block, Table):
            add_table(document, block.headers, block.rows, block.widths, block.font_size)
        elif isinstance(block, Figure):
            figure_number += 1
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            for index, image in enumerate(block.paths):
                if index:
                    p.add_run("   ")
                p.add_run().add_picture(str(image), width=Inches(6.75 * block.width / 100))
            add_caption_docx(document, f"Figura {figure_number}  {block.caption}")
        elif isinstance(block, Code):
            add_code_docx(document, block.text, block.caption)
    document.save(path)


def main() -> None:
    doc = build_content()
    html_dir = ROOT / "output" / "html"
    pdf_dir = ROOT / "output" / "pdf"
    docx_dir = ROOT / "output" / "docx"
    for directory in (html_dir, pdf_dir, docx_dir):
        directory.mkdir(parents=True, exist_ok=True)
    html_path = html_dir / f"{NAME}.html"
    html_path.write_text(render_html(doc), encoding="utf-8")
    render_pdf(html_path, pdf_dir / f"{NAME}.pdf")
    render_docx(doc, docx_dir / f"{NAME}.docx")
    for path in (html_path, pdf_dir / f"{NAME}.pdf", docx_dir / f"{NAME}.docx"):
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
